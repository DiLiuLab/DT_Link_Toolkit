#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
strand_passage_guiV4_3.py  (V4.3)
=================================

Interactive strand-passage explorer with component-colour preservation.

What is new in V4.3
-------------------
* DT codes may carry a FLIP VECTOR: ``--dt``, the GUI's DT field, a drawing
  session's ``texts.dt`` and ``--demo`` accept
  ``DT: [(10, 6, 2, 4, 12, 8)], [0,1,0,0,0,1]`` (also ', 010001', '(0,1,..)',
  T/F, str() of spherogram's DT_code(flips=True)).  A signed DT code alone
  does not fix the diagram -- the code above is the granny knot, its mirror
  or the SQUARE knot -- and the vector names one: bit k belongs to the k-th DT
  entry (odd label 2k+1); walking along the visit with the smaller label,
  bit 0 = the other strand crosses from your right to your left (spherogram /
  SnapPy convention); complementing every bit is the mirror image.  An
  illegal vector (wrong length, or no diagram on the sphere) is refused.
  With a vector the realisation is kept through every step: the diagram
  carries its crossing signs (``link_engine_v4_1``), a passage keeps the
  vector (a crossing change does not touch the shadow), Simplify and every
  SnapPy step build links with ``DTcodec(comps, flips)`` and export verified
  (DT, flips) pairs, and the Jones polynomials are those of the given
  realisation.  Measured (``--nongui`` under Sage): on
  ``DT: [(10, 6, 2, 4, 12, 8)], [0,1,0,0,0,1]`` (the square knot) the
  passages at DT entries 0, 4 and 5 give the LEFT-handed trefoil and those at
  1-3 the right-handed one; V4.2, which read the bare code as the granny, gave
  the right-handed trefoil for all six.
* ``--nongui`` with a vector: the DT columns of the passage sheets carry the
  vector ('DT_code_before_simplify' etc. in the flips spelling) and new
  columns ``dt_flip_vector_before_simplify`` / ``_after_simplify`` /
  ``_chosen`` (named apart from ``flipped_crossing`` / ``arriving_flips``,
  which mean CROSSING CHANGES); ``card_index`` gains ``dt_flip_vector``;
  ``run_info`` gains ``input_dt_flip_vector``; the overview cards are drawn
  and captioned with their vectors, and the header's start code is no longer
  cut at 66 characters (it is wrapped instead).  With ``--negative-even under``
  the code is first rewritten in spherogram's convention (entries negated;
  the vector, a property of the shadow, is unchanged).
* Fix: ``advance`` rebuilt the direct after-passage diagram from
  ``to_signed_dt()`` -- always in the 'over' convention -- with the caller's
  ``negative_even``, so with ``--negative-even under`` every crossing of the
  direct diagram was switched relative to the passage actually made (the
  engine's SnapPy rebuild had the same bug).  Both now read it as 'over'.
* Fix (same root cause): ``--nongui --negative-even under`` handed the
  user's BARE code to SnapPy as is, i.e. in the 'over' convention.  The link
  SnapPy analysed was nevertheless the right one -- spherogram picks the
  realisation of ``-X`` that is ``X``'s seen from the other side, so a bare
  code and its negation read as the same link (measured: writhe, Jones and
  linking matrix equal on 59 codes, the corpus and all its crossing
  changes) -- but the workbook mixed conventions ('under' direct codes next
  to SnapPy's 'over' codes) and the overview drew every card, SnapPy's
  codes included, in 'under': each SnapPy-simplified card showed the
  diagram with every crossing switched.  The bare code is now rewritten in
  the 'over' convention first, exactly as a flips code already was: an
  [info] line, the run_info row ``negative_even_note``, every code and
  ``flipped_crossing`` label in 'over', every card drawn as its code.
  'over' runs are unchanged.
* A bare code (no vector) gives exactly the V4.2 results otherwise -- unless
  the composite gate flags it (2026-10-19).  A bare code fixes a connected
  sum or split union only up to reflecting its parts independently, and
  every bare re-read lets spherogram choose again: measured on the bare
  granny, ``--nongui`` gave every passage the right-handed trefoil's Jones
  polynomial for the square-reading code (3 of the 6 are left-handed), a
  non-minimal granny input was re-rooted onto the square knot in 6 of 7
  workbooks, and Simplify showed the square knot for the granny 10 times in
  10.  A bare root that ``link_engine.composite_danger`` flags (two or more
  independently reflectable parts knotted or linked; never a prime link) is
  now followed as spherogram reads it (``link_engine.spherogram_reading``:
  the realisation ``snappy.Link(bare)`` builds) through the flips machinery
  above -- the GUI's root, ``--nongui`` and ``--demo``, with one [info] line
  (``--nongui``: the flips columns and ``run_info``'s
  ``dt_flip_vector_note``).  If spherogram cannot read the code a [warn]
  says so and the bare run goes ahead as before.  Simplify and every
  passage on a bare diagram go through the engine's gated
  ``snappy_global_simplification``, which keeps a flagged code's
  realisation the same way.  Unflagged runs are unchanged.
* The carried realisation (2026-10-19, second pass of the gate).  Every
  bare root now CARRIES its realisation -- spherogram's reading, or
  dt_converter's mended reading for a flagged root spherogram reads
  illegally (the split 'DT: [(4,6,2),(10,12,8)]' ran into "Link isn't
  planar"; it now runs as 001001) -- and a passage keeps the vector.  It
  is written out (the flips spelling, the ``dt_flip_vector_*`` and card
  ``dt_flip_vector`` columns, '' on the other rows of that sheet) ONLY on
  codes the composite gate flags, and it is followed from the first flagged
  code on.  ``--nongui``'s bare re-reads all keep it there: the source
  simplification (a root re-rooted onto a flagged code follows the
  simplified link's own vector: the 10-crossing prime diagram of the square
  knot 'DT: [(14, -8, 10, -2, 16, 18, 20, -6, -4, 12)]' used to give six
  right-handed trefoils in the Jones column, now three of each hand), each
  passage (a flagged direct code follows the carried vector), SnapPy's
  flagged simplified code (its own vector: 8_19's passage at -2 is the
  granny in every run), the chosen code, the second-pass roots and the
  reconciliation.  In the GUI a passage whose direct code is flagged follows
  the carried vector, not spherogram's reading of the direct code.  A flagged
  bare root's run (round 2: every code with its vector) now writes vectors
  only on its flagged codes.  A run that meets no flagged code is unchanged.
* (2026-10-19, round 4) No bare decode hangs any more: importing the engine
  installs dt_converter's guard on spherogram's flips-free DT decoder, which
  raises DecoderLoop where it would loop for ever ('DT: [(8,6,2,4)]', a
  trefoil with a kink: ``--nongui`` and ``--demo`` never returned, and the
  GUI's Load froze).  ``--nongui`` on a bare root spherogram cannot build
  (DecoderLoop, 'not enough values to unpack' on a kink, 'PD code isn't
  consistent' or "Link isn't planar" on split codes) used to end in a
  traceback; it now follows a legal reading of it -- spherogram's own where
  only its string parser fails, dt_converter's (the smallest flip vector
  spherogram builds), or a legal flip vector spherogram builds read off a
  planar embedding (``link_engine.buildable_legal_flips``) -- with one [info]
  line, exit 0 (the root is written with that vector, every other code the
  composite gate does not flag still bare); a root with no reading at all, or
  with no crossings, is one [error]
  line, exit 2.  A FLAGGED root spherogram's decoder cannot read is followed
  the same way in the GUI, ``--demo`` and ``--nongui`` (it was read bare
  with a [warn]).  A bare code that no diagram on the sphere has
  ('DT: [(4,6,8,10,2)]', 'DT: [(8,10,12),(2,4,6)]') is refused in one line
  by ``--nongui`` and ``--demo`` and in the error window of the GUI's Load;
  each ended in a traceback.  Codes spherogram reads run exactly as before.
* (2026-10-20, round 4b) The [info] lines and the ``dt_flip_vector_note`` of
  a followed unflagged root say that the root is written with its vector
  (they said only flagged codes were); the one-line no-diagram refusal of
  ``--nongui`` / ``--demo`` quotes the code without its '# comment'; every
  bare ``snappy.Link`` read goes through ``_bare_snappy_link``, which retries
  the decoder guard's installation if it failed at import.
* (2026-10-20, round 4d) ``--nongui`` and ``--demo`` on a BARE code the
  parser cannot read -- an alphabetical 'DT[daddcab]', a prefix-less
  '8,6,2,4' -- print one [error] line in the scorer's wording
  (``_unparseable_text``: what this tool reads, and that dt_converter.py
  converts other spellings), exit 2; they ended in a traceback, with a
  per-process object address in it.  ``--nongui`` on an unflagged root
  with an EMPTY component that spherogram cannot read ('DT: [(8,6,2,4),
  ()]': dt_converter refuses the empty component) no longer advises a flip
  vector -- spherogram builds no realisation with a trailing empty
  component -- but names the component to drop (without it the run
  follows a reading, exit 0).  Codes that ran run exactly as before.
* Imports ``draw_dt_original_labelsV5_9.py`` and ``link_engine_v4_1.py``.

What is new in V4.2
-------------------
* Overview SVGs no longer declare their fonts with the CSS ``font:`` shorthand.
  Matplotlib's ``svg.fonttype: none`` writes every text run as
  ``font: 700 8.6px 'Arial'``; the shorthand resets every property it does not
  name, and a consumer that implements it only partially (Adobe Illustrator
  being the one that bites here) can drop the size or weight and re-lay-out the
  string.  Because each run carries a single anchor and each background box is a
  separate path sized from Matplotlib's own measurements, any such re-layout
  slides the text out of its box.  ``_rewrite_svg_font_declarations`` now
  rewrites each declaration as explicit longhand after export:

      font: 700 8.6px 'Arial'
      -> font-family: Arial, 'Helvetica Neue', Helvetica, sans-serif;
         font-size: 8.6px; font-weight: 700

  Geometry, text content and anchors are untouched -- WebKit renders the before
  and after files to byte-identical rasters -- so this costs correct consumers
  nothing.
* The rewrite also adds a fallback family stack.  Previously the only family
  named was ``Arial`` with no generic fallback, so a machine without Arial got a
  silent substitution *and* the metric mismatch that follows from it.
* Text stays fully editable in Illustrator; this is deliberately NOT a switch to
  ``svg.fonttype: 'path'``, which would guarantee fidelity by converting text to
  outlines but undo the editable-text work from V3.7.

What is new in V4.1
-------------------
* ``--nongui`` workbook rows are now **shaded by the overview-SVG card** they
  are drawn as.  Every passage row lands on exactly one merged structure, and
  that structure is one card in ``<out>_overview.svg``; rows that merge into the
  same card share one background shade, so the table and the figure can be read
  together at a glance.
* A new ``svg_card_id`` column carries the same association as text, so it
  survives CSV export, sorting, and colour-blind viewing.
* A new ``card_index`` sheet is the legend: one shaded row per card, with its
  depth, DT code, crossing/component counts, Jones polynomial, outcome tag, the
  passages arriving at it, and the sheets whose rows it holds.
* The final per-step reconciliation now runs **before** the workbook is written
  instead of after it.  Previously the reconciliation could merge or further
  simplify structures once the spreadsheet was already on disk, so a workbook
  could describe a structure that no card in the figure actually showed.  The
  two outputs are now built from the same final graph.

What is new in V4.0
-------------------
* Added a per-window "Simplify" button.  It simplifies the currently displayed
  diagram (original or after-passage) with the active SnapPy/backtrack settings
  and refreshes the drawing/properties panel, including Jones data.
* The GUI now has one combined crossing-label field.  Assignment-style text is
  detected as a crossing map; other text is detected as crossing order.
* The strand-passage default drawing layout is now ``shaped-tutte`` with an
  ellipse boundary and manual Tutte aspect 1.0.
* ``--drawing-session`` and the GUI "Load drawing session" button load drawing
  settings from JSON sessions saved by ``draw_dt_original_labelsV5_4.py``.  GUI,
  demo, and ``--nongui`` overview drawings reuse those settings for following
  strand-passage diagrams.
* The top control strip has light-blue "?" help buttons for SnapPy/global and
  backtrack simplify settings.

What is new in V3.8
-------------------
* ``--nongui`` overview SVG layout now gives editable text more breathing room:
  overview text is set to Arial before figure construction, card/footer panels
  are slightly roomier, and text-background boxes have larger padding so the
  final SVG better matches the Matplotlib-rendered view in Illustrator.
* Drawing/model layer is now ``draw_dt_original_labelsV5_4.py``; its editable
  SVG text boxes/circles also use Arial and larger padding by default.

What is new in V3.7
-------------------
* ``--nongui`` overview SVGs now keep text as editable SVG text instead of
  expanding labels to outline paths.  Overview text artists are forced to Arial
  before export to reduce Adobe Illustrator font-substitution warnings while
  keeping labels selectable/editable.

What is new in V3.6
-------------------
* ``--nongui`` second-pass continuation now uses the criterion
  ``new_components > 2`` (plus a usable chosen DT code), and prints that
  criterion in the CLI output.
* ``--nongui`` workbooks include a ``run_info`` sheet with software versions,
  runtime details, command/arguments, output paths, key parameters, and
  continuation counts for the current result.

What is new in V3.5
-------------------
* ``--nongui`` now sends only one representative of each merged first-step
  structure into the second strand-passage pass.  The first-step graph is
  reconciled/merged before continuation, matching the overview cards more
  closely and avoiding duplicate second-pass computations from equivalent
  first-pass diagrams.
* Second-pass workbook sheets are named after the merged first-step node, and
  second-pass rows record the first-step passages that merged into that node
  plus the representative passage used for continuation.

What is new in V3.4
-------------------
* GitHub-ready packaging: public README, MIT license, dependency notes, and
  executable-script guidance.
* Optional Tk window/task-menu icon loaded from ``assets/strand_passage_icon.png``
  when present.  Missing or unsupported icon assets are ignored, so the scripts
  still run from a plain source checkout.
* Drawing/model layer is now ``draw_dt_original_labelsV5_4.py``.

What is new in V3.3
-------------------
* Optional backtrack-assisted SnapPy simplification, in BOTH the GUI and
  ``--nongui``.  SnapPy's ``simplify('global')`` is greedy and can stop at a
  non-minimal diagram (e.g. 12 crossings for a link whose minimum is 10), which
  affects the drawn diagrams and the DT-code choice everywhere.  Enabling
  backtrack repeatedly complicates then re-simplifies to escape such plateaus,
  keeping the fewest-crossing diagram.  Controlled by the number of rounds and
  the backtrack step size; OFF by default (rounds = 0) so nothing changes unless
  you ask for it.  CLI: ``--backtrack --backtrack-rounds N --backtrack-steps K``;
  GUI: the "Backtrack simplify" checkbox plus rounds/steps fields.

What is new in V3.2
-------------------
* Drawing/model layer is now ``draw_dt_original_labelsV5_4.py``, and 2-D links
  are drawn with that helper's own
  DEFAULT settings (default layout, top-to-bottom orientation, and
  false-crossing visualization).
* DT-code choice rule, applied everywhere (GUI and ``--nongui`` spreadsheet):
  after a passage, if the *direct after-passage DT code* has MORE crossings than
  the SnapPy ``simplify('global')`` code, the SnapPy-simplified code is used;
  otherwise (fewer-or-equal, i.e. a tie) the direct after-passage DT code is
  kept.
* DT traversal labels stay visible in the view even after a SnapPy
  simplification (the chosen diagram is drawn from its own DT code).
* The mouse cursor turns into a pointing hand when it is near a crossing where a
  strand passage can be performed.
* Optimised window layout: a resizable diagram/properties split, scrollable
  properties, a status/hint bar, and Save PNG / Save SVG that reproduce exactly
  what is shown on screen.
* 2-D views match the helper's default look, including crossing IDs coloured by
  the over-strand's component colour.
* The pointing-hand cursor uses an adaptive radius (a little under half the
  nearest-crossing gap), so it only appears when the cursor is genuinely on a
  crossing.
* A warning is printed when not running under Sage (Jones polynomials, and the
  SnapPy invariant colour-matching that relies on them, are unavailable).
* --nongui also writes a large, tidy overview SVG (<name>_overview.svg) of every
  resulting structure across the two passage steps: colour-coded cards with the
  2-D view, DT code, crossing/component counts, Jones, a component colour key,
  and outcome tags; straight arrows labelled with the flipped crossing show the
  operation order; topologically identical structures are merged into one card.

Non-interactive spreadsheet (behaves like the old strand_pass_sage.py):
    sage -python strand_passage_guiV4_3.py --nongui \
        --dt "DT: [(-8,-12,16),(-24,-22,-28,-26),(-10,-14,-2),(-20,-6,-18,-4)]" \
        --out strand_passage_results.xlsx

Interactive run:
    sage -python strand_passage_guiV4_3.py                 # SnapPy enabled
    sage -python strand_passage_guiV4_3.py --dt "DT: [(4,6,2)]"
    sage -python strand_passage_guiV4_3.py --dt "DT: [(10, 6, 2, 4, 12, 8)], [0,1,0,0,0,1]"
    python3 strand_passage_guiV4_3.py --gui-backend agg    # if TkAgg won't load

Headless cascade figure (no display needed):
    python3 strand_passage_guiV4_3.py --dt "DT: [(4,6,2)]" --demo 2 1 --out chain.png
"""

from __future__ import annotations

import argparse
import ast
import copy
from datetime import datetime
import json
import os
import re
import shlex
import sys
from typing import Any, Dict, List, Optional, Tuple

import numpy as np

# The drawing helper imports matplotlib.pyplot at import time; pin a safe global
# backend so Sage/macOS never fails before the GUI's own backend fallback runs.
os.environ.setdefault("MPLBACKEND", "Agg")

import draw_dt_original_labelsV5_9 as D          # noqa: E402
import link_engine_v4_1 as E                       # noqa: E402

TAB10_NAMES = ["blue", "orange", "green", "red", "purple",
               "brown", "pink", "gray", "olive", "cyan"]
DEFAULT_DT = "DT: [(-8,-12,16),(-24,-22,-28,-26),(-10,-14,-2),(-20,-6,-18,-4)]"
VERSION = "4.3"
OVERVIEW_FONT_FAMILY = "Arial"
OVERVIEW_TEXT_BOX_PAD = 0.42
OVERVIEW_DIAGRAM_DT_LABEL_BOX_PAD = 0.22
OVERVIEW_DIAGRAM_CROSSING_ID_BOX_PAD = 0.28
NONGUI_SECOND_PASS_CRITERION = (
    "first-step new_components > 2 and DT_code_chosen is available")
DEFAULT_BACKTRACK_ROUNDS = getattr(E, "DEFAULT_BACKTRACK_ROUNDS", 200)
DEFAULT_BACKTRACK_STEPS = getattr(E, "DEFAULT_BACKTRACK_STEPS", 30)
DEFAULT_DRAWING_OPTIONS = E.default_drawing_options()

# Tile size used by link_engine's render() when packing blocks.
TILE = 2.4
# Fallback click/hover radius (data units) for a diagram with a single crossing,
# where there is no neighbour distance to calibrate against.
SINGLE_CROSSING_RADIUS = 0.28


# --------------------------------------------------------------------------- #
#  DT + flip vector (V4.3).  The block between the BEGIN and END markers is the
#  toolkit's shared helper, byte-identical in every tool that reads DT text.
#  'DT: [(10, 6, 2, 4, 12, 8)], [0,1,0,0,0,1]': bit k belongs to the k-th DT
#  entry (odd label 2k+1); walk along the visit with the SMALLER label -- bit 0
#  when the other strand crosses from your right to your left.  Complementing
#  every bit reflects the page (the mirror image).  The components themselves
#  are read by link_engine's parse_dt_flips.
# --------------------------------------------------------------------------- #
# BEGIN dt-flips helpers (keep identical in every tool)
_DT_FLIPS_ERROR = ("could not read %r after the DT code as a flip vector "
                   "(one 0/1 per crossing, e.g. 'DT: [(4, 6, 2)], [0,0,1]')")


def split_dt_flips(text):
    """Split 'DT: [...], [0,1,...]' into ('DT: [...]', [False, True, ...]).

    Text with nothing after the DT code's closing bracket comes back
    UNCHANGED with None, so a bare code takes exactly its old path; so does
    an old multi-list spelling whose later lists hold more components
    ('DT: [(-6,-8)], [(-4,-2)]').  Also read: ', 010001' / '; 010001' / no
    separator, '(0,1,..)', spaces instead of commas, T/F and True/False, and
    the tuple wrapper 'DT: ([(..)], [0,1,..])' (str() of spherogram's
    DT_code(flips=True)).  A vector has one bit per entry of the code's list.
    A trailing '# comment' is dropped (a bare code with one comes back
    without it).  Any other text after the code raises ValueError -- it is
    never ignored."""
    s = str(text)
    t = s.rstrip()
    i = t.find("[")
    if i < 0:
        return s, None
    depth, j = 0, -1
    for j in range(i, len(t)):
        if t[j] == "[":
            depth += 1
        elif t[j] == "]":
            depth -= 1
            if depth == 0:
                break
    if depth != 0:
        if t.count("[") > 1:
            raise ValueError("unbalanced brackets in %r: the code's list must "
                             "close before a flip vector, e.g. "
                             "'DT: [(4, 6, 2)], [0,0,1]'" % s[:80])
        return s, None            # unbalanced: the caller's parser reports it
    head, rest = t[:i], t[j + 1:]
    note = rest.find("#")
    base = s if note < 0 else t[:j + 1 + note].rstrip()
    rest = (rest if note < 0 else rest[:note]).strip()
    wrapped = head.rstrip().endswith("(") and rest.endswith(")")
    if wrapped:
        rest = rest[:-1].strip()
    if not rest:
        inner = t[i + 1:j]
        if "(" in inner and re.search(r"\d", re.sub(r"\([^()]*\)", "", inner)):
            raise ValueError("numbers outside the parenthesised components of "
                             "%r: a flip vector goes AFTER the code's closing "
                             "bracket, e.g. 'DT: [(4, 6, 2)], [0,0,1]'" % s[:80])
        return base, None
    code = (head.rstrip()[:-1] if wrapped else head) + t[i:j + 1]
    body = rest[1:].strip() if rest[0] in ",;" else rest
    if body and body[0] + body[-1] in ("[]", "()"):
        body = body[1:-1].strip()
    tokens = [x for x in re.split(r"[\s,;]+", body) if x]
    bits = None
    if len(tokens) == 1 and re.fullmatch(r"[01]+|[TFtf]+", tokens[0]):
        bits = [ch in "1Tt" for ch in tokens[0]]
    elif tokens and all(x.lower() in ("0", "1", "t", "f", "true", "false")
                        for x in tokens):
        bits = [x.lower() in ("1", "t", "true") for x in tokens]
    labels = [int(v) for v in re.findall(r"\d+", t[i + 1:j])]
    n = len(labels)
    if bits is not None and len(bits) == n:
        return code, bits
    labels += [int(v) for v in re.findall(r"\d+", rest)]
    if (len(labels) > n and sorted(labels) == list(range(2, 2 * len(labels) + 1, 2))
            and not re.search(r"[^\s,;|()\[\]\d\-\u2212\u2012\u2013\u2014\ufe63\uff0d]", rest)):
        return base, None         # more components: an old multi-list spelling
    if bits is not None:
        raise ValueError("the flip vector has %d entries but the code has %d "
                         "crossings" % (len(bits), n))
    raise ValueError(_DT_FLIPS_ERROR % rest)


def format_dt_flips(dt_text, flips):
    """'DT: [...]' + flips -> 'DT: [...], [0,1,...]' (dt_converter's spelling);
    the bare text unchanged when flips is None."""
    if flips is None:
        return dt_text
    return "%s, [%s]" % (str(dt_text).rstrip(),
                         ",".join("1" if f else "0" for f in flips))


def check_flips(comps, flips):
    """Raise ValueError unless ``flips`` is a legal vector for the signed DT
    code ``comps`` (a list of tuples of even integers): one bit per crossing,
    and the diagram it describes lies on the sphere (every connected piece
    has genus 0: faces == crossings + 2 x pieces).  Pure Python; accepts the
    kink vectors spherogram's DTcodec cannot build."""
    entries = [int(e) for c in comps for e in c]
    n = len(entries)
    flips = [bool(f) for f in flips]
    if len(flips) != n:
        raise ValueError("the flip vector has %d entries but the code has %d "
                         "crossings" % (len(flips), n))
    cross, nxt, start = {}, {}, 1
    for c in comps:
        labels = list(range(start, start + 2 * len(c)))
        for a, b in zip(labels, labels[1:] + labels[:1]):
            nxt[a] = b
        start += 2 * len(c)
    for k, e in enumerate(entries):
        cross[2 * k + 1] = cross[abs(e)] = k
    if sorted(cross) != list(range(1, 2 * n + 1)):
        raise ValueError("not a DT code: the even labels are not a "
                         "permutation of 2..%d" % (2 * n))

    def slots(p):                 # (in slot, out slot) of label p: S0 E1 N2 W3
        k = cross[p]
        small = min(2 * k + 1, abs(entries[k]))
        if p != small:
            return 1, 3
        return (2, 0) if flips[k] else (0, 2)

    other = {}
    for p in range(1, 2 * n + 1):
        q = nxt[p]
        a, b = (cross[p], slots(p)[1]), (cross[q], slots(q)[0])
        other[a], other[b] = b, a
    seen, faces = set(), 0
    for d in other:
        if d in seen:
            continue
        faces += 1
        while d not in seen:
            seen.add(d)
            k, sl = other[d]
            d = (k, (sl + 1) % 4)
    parent = list(range(n))

    def root(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x
    for p in range(1, 2 * n + 1):
        parent[root(cross[p])] = root(cross[nxt[p]])
    pieces = len({root(k) for k in range(n)})
    if faces != n + 2 * pieces:
        raise ValueError(
            "the flip vector %s is not a legal one for this DT code: it "
            "describes no diagram on the sphere (%d faces; a planar diagram "
            "with %d crossing(s) in %d piece(s) has %d)"
            % ("".join("1" if f else "0" for f in flips), faces, n, pieces,
               n + 2 * pieces))


def flips_crossing_signs(comps, flips):
    """Crossing signs (+1 / -1, in DT entry order) of the diagram (comps,
    flips) in spherogram's convention, where a NEGATIVE entry means the even
    visit is the over-crossing: negative exactly when
    flip XOR (even visit over) XOR (even label < odd label)."""
    out = []
    for k, e in enumerate(int(e) for c in comps for e in c):
        neg = bool(flips[k]) ^ (e < 0) ^ (abs(e) < 2 * k + 1)
        out.append(-1 if neg else 1)
    return out
# END dt-flips helpers


def _looks_like_crossing_map(text):
    """Heuristic for the combined crossing-label field."""
    return bool(re.search(r"[cC]?\d+\s*[:=]\s*\d+", text or ""))


def _split_crossing_display_input(crossing_input=None, crossing_order=None,
                                  crossing_map=None):
    """Return ``(crossing_order, crossing_map, detected_kind)``.

    ``crossing_input`` is the V4.0 combined field.  Assignment-style text such as
    ``c1=1,c7=3`` is treated as a map; otherwise the text is treated as an order
    list.  The old explicit CLI/session fields are still accepted for backward
    compatibility.
    """
    combined = (crossing_input or "").strip()
    order = (crossing_order or "").strip()
    cmap = (crossing_map or "").strip()
    if combined:
        if order or cmap:
            raise ValueError(
                "Use either the combined crossing field or the legacy "
                "crossing-order/crossing-map fields, not both.")
        if _looks_like_crossing_map(combined):
            return "", combined, "crossing-map"
        return combined, "", "crossing-order"
    if order and cmap:
        raise ValueError("Use either crossing-order or crossing-map, not both.")
    if cmap:
        return "", cmap, "crossing-map"
    if order:
        return order, "", "crossing-order"
    return "", "", ""


def _apply_crossing_display_options(diagram, crossing_order=None,
                                    crossing_map=None, strict=True,
                                    crossing_input=None):
    """Attach custom displayed crossing IDs to a diagram.

    The labels are display-only; strand-passage operations still use the
    diagram's internal crossing ids.
    """
    try:
        crossing_order, crossing_map, _kind = _split_crossing_display_input(
            crossing_input=crossing_input,
            crossing_order=crossing_order,
            crossing_map=crossing_map,
        )
    except Exception:
        if strict:
            raise
        crossing_order, crossing_map = "", ""
    if not crossing_order and not crossing_map:
        if hasattr(diagram, "crossing_display_ids"):
            delattr(diagram, "crossing_display_ids")
        return diagram
    model = getattr(diagram, "_dt_model", None)
    if model is None:
        if strict:
            raise ValueError("custom crossing IDs need a diagram built from a DT code")
        return diagram
    try:
        ids = D.resolve_crossing_ids(
            model,
            crossing_order=crossing_order or None,
            crossing_map=crossing_map or None,
        )
    except Exception:
        if strict:
            raise
        return diagram
    if len(ids) != len(diagram.crossings()):
        if strict:
            raise ValueError(
                "custom crossing IDs describe %d crossings, but this diagram has %d"
                % (len(ids), len(diagram.crossings())))
        return diagram
    diagram.crossing_display_ids = list(ids)
    return diagram


def _inherit_crossing_display_ids(source, target):
    """Carry display labels across a passage when the crossing count still fits."""
    ids = getattr(source, "crossing_display_ids", None)
    if ids and len(ids) == len(target.crossings()):
        target.crossing_display_ids = list(ids)
    return target


def _crossing_display_label(diagram, crossing_id):
    ids = getattr(diagram, "crossing_display_ids", None)
    if ids and 0 <= int(crossing_id) < len(ids):
        return str(ids[int(crossing_id)])
    return "c%d" % (int(crossing_id) + 1)


def _hit_radius(targets):
    """Adaptive click/hover radius so the pointing hand only appears when the
    cursor is genuinely ON a crossing.

    With two or more crossings the radius is a little under half the smallest
    gap between crossings, so hover/click zones never overlap and never reach
    far into empty space.  With a single crossing a small fixed radius is used.
    """
    pts = [np.asarray(xy, float) for xy in targets.values()]
    if len(pts) < 2:
        return SINGLE_CROSSING_RADIUS
    best = float("inf")
    for i in range(len(pts)):
        for j in range(i + 1, len(pts)):
            best = min(best, float(np.hypot(*(pts[i] - pts[j]))))
    if not np.isfinite(best):
        return SINGLE_CROSSING_RADIUS
    return max(0.10, min(SINGLE_CROSSING_RADIUS, 0.42 * best))


def _running_under_sage():
    """True when a Sage runtime is importable (``sage -python``)."""
    try:
        import sage  # type: ignore  # noqa: F401
        return True
    except Exception:  # noqa: BLE001
        return False


def warn_if_no_sage():
    """Warn that Jones polynomials need Sage/SnapPy when running plain python."""
    if _running_under_sage():
        return
    sys.stderr.write(
        "[warning] Not running under Sage: Jones polynomials (and the SnapPy "
        "invariant colour-matching that relies on them) cannot be computed. "
        "For full functionality run this with 'sage -python "
        "strand_passage_guiV4_3.py ...'.\n")


_DRAWING_STRING_KEYS = {
    "layout", "y_direction", "tutte_shape",
}
_DRAWING_FLOAT_KEYS = {
    "rotate", "font_size", "crossing_id_font_size", "line_width", "gap_frac",
    "tutte_aspect", "tutte_corner_radius", "tutte_decompress",
    "tutte_com_expand", "tutte_orient", "hole_ratio", "ring_tilt", "min_sep",
}
_DRAWING_BOOL_KEYS = {
    "show_crossing_ids", "color_crossing_ids_by_overstrand", "hide_labels",
    "no_arrows", "tutte_auto_aspect", "tutte_auto_orient",
    "show_tutte_outline", "show_tutte_pca", "hole_swap", "invert_ring",
}


def default_drawing_options():
    """Drawing defaults shared by GUI, demo, and nongui overview rendering."""
    return E.default_drawing_options()


def _coerce_session_float(key, value):
    try:
        return float(value)
    except Exception:  # noqa: BLE001
        return DEFAULT_DRAWING_OPTIONS[key]


def _coerce_session_bool(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    text = str(value).strip().lower()
    if text in ("1", "true", "yes", "on"):
        return True
    if text in ("0", "false", "no", "off"):
        return False
    return bool(value)


def normalize_drawing_options(options=None):
    """Merge a partial drawing-options dict onto the V4.0 defaults."""
    out = default_drawing_options()
    for key, value in dict(options or {}).items():
        if key not in out or value is None:
            continue
        if key in _DRAWING_FLOAT_KEYS:
            out[key] = _coerce_session_float(key, value)
        elif key in _DRAWING_BOOL_KEYS:
            out[key] = _coerce_session_bool(value)
        else:
            out[key] = str(value)
    return out


def load_drawing_session(path):
    """Read a session JSON written by ``draw_dt_original_labels*.py`` (V5.9).

    Returns a small dict with drawing options plus optional DT/crossing-label
    settings.  Unknown helper-only 3-D fields are ignored.
    """
    with open(path, "r", encoding="utf-8") as fh:
        data = json.load(fh)
    strings = data.get("strings") or {}
    bools = data.get("bools") or {}
    texts = data.get("texts") or {}

    drawing = {}
    for key, value in strings.items():
        if key in _DRAWING_STRING_KEYS or key in _DRAWING_FLOAT_KEYS:
            drawing[key] = value
    for key, value in bools.items():
        if key in _DRAWING_BOOL_KEYS:
            drawing[key] = value

    crossing_order = str(texts.get("crossing_order") or "").strip()
    crossing_map = str(texts.get("crossing_map") or "").strip()
    crossing_input = crossing_map or crossing_order
    negative_even = str(strings.get("negative_even") or "").strip() or None
    if negative_even not in ("over", "under"):
        negative_even = None

    return {
        "path": os.path.abspath(path),
        "script": data.get("script", ""),
        "version": data.get("version", ""),
        "drawing_options": normalize_drawing_options(drawing),
        "dt": str(texts.get("dt") or "").strip(),
        "crossing_input": crossing_input,
        "crossing_order": crossing_order,
        "crossing_map": crossing_map,
        "negative_even": negative_even,
    }


# --------------------------------------------------------------------------- #
#  Core (GUI-independent): one strand passage + DT-code choice
# --------------------------------------------------------------------------- #
def advance(diagram, crossing_id, negative_even="over", use_snappy_global=True,
            backtrack_rounds=0, backtrack_steps=20):
    """
    Perform a strand passage at ``crossing_id`` on ``diagram`` and return

        (result_diagram, note, snap)

    Choice rule (V3.2)
    ------------------
    Let the *direct after-passage DT code* be the diagram obtained by flipping the
    crossing (no simplification); it has the same crossing count as ``diagram``.
    If SnapPy's ``simplify('global')`` yields STRICTLY FEWER crossings, that
    SnapPy-simplified diagram is chosen; otherwise the direct after-passage
    diagram is kept.  Either way the chosen diagram is built from its own DT code
    so its original DT traversal labels remain visible.

    ``backtrack_rounds`` > 0 enables backtrack-assisted SnapPy simplification
    (V3.4), so the SnapPy-simplified crossing count -- and hence the choice -- is
    based on a diagram that escaped ``simplify('global')`` plateaus.

    ``diagram`` itself is never mutated.

    V4.3: ``to_signed_dt()`` writes the 'over' convention (a negative even
    label = the even visit is over) whatever ``negative_even`` the root was
    read with, so the direct diagram is rebuilt in 'over' (rebuilding it with
    'under' switched every crossing).  A diagram that carries crossing signs
    (its root had a flip vector) keeps its realisation: the passage keeps the
    vector, and the direct diagram and SnapPy get the recomputed
    (DT, flips) pair.  Round 3 (2026-10-19): a bare diagram that CARRIES a
    realisation (its bare root's reading) keeps it the same way -- the direct
    diagram carries the vector, and SnapPy simplifies that realisation when
    the composite gate flags the direct code or a flagged code was met on the
    way -- but writes it out only where the gate flags a code; otherwise
    exactly as before.
    """
    if crossing_id not in diagram.crossings():
        raise ValueError("crossing c%d is not present" % (crossing_id + 1))

    before = len(diagram.crossings())
    source_kind = getattr(diagram, "display_source", "continuity")

    working = diagram.copy()
    working.crossing_change(crossing_id)
    flips_known = getattr(working, "crossing_signs", None) is not None
    carried = None
    if flips_known:
        comps_before, flips_before, dt_ok = working.signed_dt_flips()
        dt_before = (E.dt_flips_to_string(comps_before, flips_before)
                     if dt_ok else None)
    else:
        dt_before, dt_ok = working.to_signed_dt()
        # V4.3 (round 3): the realisation a bare run carries -- the passage
        # keeps the vector; SnapPy gets the bare code and the diagram, and
        # follows the vector where the code is flagged (or the carried
        # realisation is active)
        if dt_ok and getattr(working, "carry_signs", None) is not None:
            _cc, _cf, _cok = working.carried_dt_flips()
            if _cok and _cf is not None:
                carried = _cf

    # Build the direct after-passage diagram from its DT so DT labels show and
    # component colours keep their original identity.
    direct = None
    if dt_ok:
        try:
            if flips_known:
                direct = E.Diagram.from_dt_code(
                    comps_before, negative_even="over",
                    component_colors=list(working.component_colors),
                    flips=flips_before)
            elif carried is not None:
                direct_code = D.parse_dt(dt_before)
                direct = E.Diagram.from_dt_code(
                    direct_code, negative_even="over",
                    component_colors=list(working.component_colors),
                    carry=carried,
                    carry_active=bool(working.carry_active or (
                        carried and E.composite_danger(direct_code)[0])))
            else:
                direct = E.Diagram.from_dt_code(
                    D.parse_dt(dt_before), negative_even="over",
                    component_colors=list(working.component_colors))
            direct.colours_tracked = getattr(diagram, "colours_tracked", True)
        except Exception:  # noqa: BLE001
            direct = None
    if direct is None:
        direct = working
    _inherit_crossing_display_ids(diagram, direct)
    direct.display_source = "direct after-passage DT"
    direct_crossings = len(direct.crossings())

    snap = None
    chosen = direct
    phrase = ""

    if use_snappy_global and dt_ok:
        snap = E.snappy_global_simplification(
            working, dt_before, negative_even=negative_even, mode="global",
            backtrack_rounds=backtrack_rounds, backtrack_steps=backtrack_steps)
        simplified_crossings = snap.get("simplified_crossings")
        snap_diagram = snap.get("diagram")
        if snap_diagram is not None:
            _inherit_crossing_display_ids(diagram, snap_diagram)
        if (snap_diagram is not None and simplified_crossings is not None
                and int(simplified_crossings) < int(direct_crossings)):
            chosen = snap_diagram
            if snap.get("colours_tracked"):
                chosen.display_source = ("SnapPy-simplified (fewer crossings; "
                                         "colours tracked)")
                phrase = ("; SnapPy-simplified chosen (%d < %d crossings; "
                          "colours tracked)"
                          % (simplified_crossings, direct_crossings))
            else:
                match = snap.get("match")
                reason = match.message if match is not None else "ambiguous match"
                chosen.display_source = ("SnapPy-simplified (fewer crossings; "
                                         "colours NOT tracked)")
                phrase = ("; SnapPy-simplified chosen (%d < %d crossings), "
                          "COLOURS NOT TRACKED (%s)"
                          % (simplified_crossings, direct_crossings, reason))
        else:
            chosen = direct
            if snap_diagram is not None and simplified_crossings is not None:
                phrase = ("; kept direct after-passage DT (%d crossings; "
                          "SnapPy gave %d, not fewer)"
                          % (direct_crossings, simplified_crossings))
            else:
                reason = (snap.get("error")
                          or (snap.get("match").message if snap.get("match") else None)
                          or snap.get("simplified_dt_error")
                          or "SnapPy produced no usable DT")
                phrase = ("; kept direct after-passage DT (SnapPy unavailable: %s)"
                          % reason)
    elif use_snappy_global and not dt_ok:
        snap = {"available": False, "source_dt": None,
                "error": "diagram has no legal DT export for SnapPy"}
        chosen = direct
        phrase = "; kept direct after-passage structure (no legal DT for SnapPy)"
    else:
        chosen = direct
        phrase = "; SnapPy global disabled: showing direct after-passage DT"

    chosen.last_snap_result = snap
    after = len(chosen.crossings())
    note = ("Passage at %s (on %s): %d -> %d crossings drawn%s."
            % (_crossing_display_label(diagram, crossing_id),
               source_kind, before, after, phrase))
    return chosen, note, snap


def simplify_current_diagram(diagram, negative_even="over", use_snappy_global=True,
                             backtrack_rounds=0, backtrack_steps=20):
    """Simplify the current diagram directly with the active SnapPy settings.

    V4.3: a diagram that carries crossing signs is exported with its flip
    vector, so SnapPy simplifies that realisation."""
    before = len(diagram.crossings())
    source_kind = getattr(diagram, "display_source", "continuity")
    if getattr(diagram, "crossing_signs", None) is not None:
        dt_now, dt_ok = diagram.to_signed_dt_flips()
    else:
        dt_now, dt_ok = diagram.to_signed_dt()
    if not use_snappy_global:
        snap = {"available": False, "source_dt": dt_now if dt_ok else None,
                "operation": "manual_simplify",
                "error": "SnapPy global is disabled in the controls"}
        note = ("Simplify current diagram (on %s): skipped; SnapPy global is "
                "disabled." % source_kind)
        diagram.last_snap_result = snap
        return diagram, note, snap
    if not dt_ok:
        snap = {"available": False, "source_dt": None,
                "operation": "manual_simplify",
                "error": "current diagram has no legal DT export for SnapPy"}
        note = ("Simplify current diagram (on %s): skipped; the current drawing "
                "has no legal DT export." % source_kind)
        diagram.last_snap_result = snap
        return diagram, note, snap

    snap = E.snappy_global_simplification(
        diagram, dt_now, negative_even=negative_even, mode="global",
        backtrack_rounds=backtrack_rounds, backtrack_steps=backtrack_steps)
    snap["operation"] = "manual_simplify"
    simplified = snap.get("diagram")
    if simplified is not None:
        _inherit_crossing_display_ids(diagram, simplified)
        chosen = simplified
        if snap.get("colours_tracked"):
            chosen.display_source = "SnapPy-simplified (manual simplify; colours tracked)"
            phrase = "SnapPy-simplified; colours tracked"
        else:
            match = snap.get("match")
            reason = match.message if match is not None else "ambiguous match"
            chosen.display_source = "SnapPy-simplified (manual simplify; colours NOT tracked)"
            phrase = "SnapPy-simplified; COLOURS NOT TRACKED (%s)" % reason
    else:
        chosen = diagram
        reason = (snap.get("error")
                  or snap.get("simplified_dt_error")
                  or "SnapPy produced no usable DT")
        phrase = "kept current diagram (%s)" % reason

    chosen.last_snap_result = snap
    after = len(chosen.crossings())
    note = ("Simplify current diagram (on %s): %d -> %d crossings drawn; %s."
            % (source_kind, before, after, phrase))
    return chosen, note, snap


def properties_text(diagram, note, snap, passage_log, use_snappy_global=True):
    d = diagram
    # V4.3: with its vector (round 3: a carried one only where the gate
    # flags the code; a bare diagram exports exactly to_signed_dt)
    dt_now, dt_ok = d.to_signed_dt_flips()
    lines = []
    lines.append(note)
    lines.append("")
    lines.append("Active drawing: %s" % getattr(d, "display_source", "continuity"))
    lines.append("Click any crossing here to open the next passage in a new window.")
    lines.append("(The cursor becomes a pointing hand over a clickable crossing.)")
    lines.append("")
    tracked = getattr(d, "colours_tracked", True)
    if tracked:
        lines.append("Components (colour = original component identity):")
        for ci in range(d.num_components()):
            xs = len(d.comps[ci])
            original_ci = d.component_colors[ci]
            nm = TAB10_NAMES[original_ci % 10]
            kind = "free unknot loop" if xs == 0 else "%d crossing visit(s)" % xs
            lines.append("   current comp %d  ->  original comp %d  [%-6s]  %s"
                         % (ci + 1, original_ci + 1, nm, kind))
    else:
        lines.append("Components (colours NOT tracked from here - default order):")
        for ci in range(d.num_components()):
            xs = len(d.comps[ci])
            nm = TAB10_NAMES[d.component_colors[ci] % 10]
            kind = "free unknot loop" if xs == 0 else "%d crossing visit(s)" % xs
            lines.append("   SnapPy comp %d  [%-6s]  %s  (original identity unknown)"
                         % (ci + 1, nm, kind))
    lines.append("")
    lines.append("Passage path: %s"
                 % (" | ".join(passage_log) if passage_log else "(original)"))
    lines.append("Crossings in this drawing: %d" % len(d.crossings()))
    lines.append("DT export of this drawing: %s"
                 % (dt_now if dt_ok else "not a legal DT traversal"))

    if snap is None:
        lines.append("")
        lines.append("SnapPy global: %s"
                     % ("not run for this window."
                        if use_snappy_global else "disabled."))
        return "\n".join(lines)

    lines.append("")
    manual_simplify = snap.get("operation") == "manual_simplify"
    if manual_simplify:
        lines.append("SnapPy global simplification / current-DT invariants:")
    else:
        lines.append("SnapPy global simplification / DT-code choice:")
    if not snap.get("available"):
        lines.append("   unavailable/skipped: %s"
                     % snap.get("error", "unknown reason"))
        return "\n".join(lines)
    if snap.get("error"):
        lines.append("   error: %s" % snap.get("error"))
        return "\n".join(lines)

    lines.append("   source crossings      : %s" % snap.get("source_crossings", "n/a"))
    if manual_simplify:
        lines.append("   input diagram         : %s crossings"
                     % snap.get("source_crossings", "n/a"))
    else:
        lines.append("   direct after-passage  : %s crossings"
                     % snap.get("source_crossings", "n/a"))
    lines.append("   SnapPy simplified     : %s crossings"
                 % snap.get("simplified_crossings", "n/a"))
    if snap.get("backtrack_rounds"):
        lines.append("   backtrack simplify    : %s rounds, %s steps"
                     % (snap.get("backtrack_rounds"), snap.get("backtrack_steps")))
    lines.append("   source components     : %s" % snap.get("source_components", "n/a"))
    lines.append("   simplified components : %s" % snap.get("simplified_components", "n/a"))
    match = snap.get("match")
    if match is not None:
        lines.append("   component matching    : %s" % match.message)
        if match.is_unique:
            mp = ", ".join("simp %d -> source %d" % (k + 1, v + 1)
                           for k, v in sorted(match.mapping.items()))
            lines.append("   matching map          : %s" % mp)
    lines.append("   drawing shown         : %s"
                 % getattr(d, "display_source", "n/a"))
    if "jones" in snap:
        lines.append("   Jones polynomial      : %s" % snap.get("jones"))
    lines.append("   source linking matrix : %s"
                 % E.matrix_for_display(snap.get("source_linking_matrix")))
    lines.append("   simplified link matrix: %s"
                 % E.matrix_for_display(snap.get("simplified_linking_matrix")))
    if snap.get("source_dt"):
        lines.append("   DT fed to SnapPy      : %s" % snap.get("source_dt"))
    if snap.get("simplified_dt"):
        lines.append("   SnapPy DT after global: %s" % snap.get("simplified_dt"))
    if snap.get("composite_note"):
        # V4.3 (2026-10-19): why a bare code is followed with a flip vector
        lines.append("   composite gate        : %s" % snap.get("composite_note"))
    if snap.get("diagram_error"):
        lines.append("   diagram warning       : %s" % snap.get("diagram_error"))
    return "\n".join(lines)


# --------------------------------------------------------------------------- #
#  Headless proof-of-cascade (no interactive backend required)
# --------------------------------------------------------------------------- #
def render_chain(dt_string, clicks, out_path, negative_even="over",
                 use_snappy_global=True, backtrack_rounds=0, backtrack_steps=20,
                 crossing_order=None, crossing_map=None, crossing_input=None,
                 drawing_options=None):
    """Simulate a click sequence and render original + each step as a panel row.

    V4.3: ``dt_string`` may carry a flip vector; the realisation is kept.  A
    bare root carries its realisation (``_bare_root_realisation``: one [info]
    line when the composite gate flags it), shown where a code is flagged."""
    import matplotlib.pyplot as plt

    comps, flips = E.parse_dt_flips(dt_string)
    carry, carry_active = None, False
    if flips is None:
        _refuse_unrealisable(comps)          # round 4: one line, not a traceback
    if flips is None:
        # round 3: the root's realisation is carried, shown where flagged
        carry, flagged, line = _bare_root_realisation(comps, negative_even)
        carry_active = bool(flagged and carry)
        if line is not None:
            print(line, file=sys.stderr if line.startswith("[warn]") else sys.stdout)
    if flips is None:
        start = E.Diagram.from_dt_code(comps, negative_even=negative_even,
                                       carry=carry, carry_active=carry_active)
    else:
        start = E.Diagram.from_dt_code(comps, negative_even=negative_even,
                                       flips=flips)
    _apply_crossing_display_options(
        start, crossing_order, crossing_map, crossing_input=crossing_input)
    start.display_source = "original"
    panels = [("original", start, None)]
    current = start
    for cid in clicks:
        if cid not in current.crossings():
            if not current.crossings():
                break
            cid = current.crossings()[0]
        nxt, note, snap = advance(current, cid, negative_even=negative_even,
                                  use_snappy_global=use_snappy_global,
                                  backtrack_rounds=backtrack_rounds,
                                  backtrack_steps=backtrack_steps)
        panels.append((note, nxt, snap))
        current = nxt

    n = len(panels)
    fig, axes = plt.subplots(1, n, figsize=(6.2 * n, 6.6))
    if n == 1:
        axes = [axes]
    for ax, (note, d, _snap) in zip(axes, panels):
        E.render(d, ax, show_crossing_ids=True, show_dt_labels=True,
                 drawing_options=drawing_options)
        ax.set_title(note, fontsize=8)
    fig.tight_layout()
    fig.savefig(out_path, dpi=140, bbox_inches="tight")
    plt.close(fig)
    return panels


# --------------------------------------------------------------------------- #
#  Non-interactive spreadsheet (behaviour of the original batch workflow,
#  extended with the V3.2 DT-code choice)
# --------------------------------------------------------------------------- #
def _dt_code_to_str(dt_code):
    """Convert a DT code (list of tuples) to SnapPy's 'DT: [...]' string."""
    return E.dt_to_string(dt_code)


def _build_label_intervals(dt_code):
    """1-based (start_label, end_label) traversal interval for each component."""
    intervals = []
    current = 1
    for comp in dt_code:
        n = len(comp)
        start = current
        end = current + 2 * n - 1
        intervals.append((start, end))
        current = end + 1
    return intervals


def _crossing_count(dt_code):
    """Number of crossings in a DT code (one crossing per label pair)."""
    return sum(len(comp) for comp in dt_code)


def _normalize_dt_code(raw):
    """Turn a SnapPy DT_code() result into a plain list of int tuples."""
    if hasattr(raw, "tolist"):
        raw = raw.tolist()
    out = []
    for comp in raw:
        if hasattr(comp, "tolist"):
            comp = comp.tolist()
        out.append(tuple(int(x) for x in comp))
    return out


def strand_passage_nongui(snappy, dt_code, orig_components=None,
                          orig_crossings=None, backtrack_rounds=0,
                          backtrack_steps=20, flips=None, carry=None,
                          carry_active=False, return_meta=False):
    """
    All single strand passages on a link, choosing the DT code per V3.2 rule.

    For each crossing, flip its sign (a strand passage) and record properties of
    the resulting link.  ``DT_code_chosen`` is the SnapPy-simplified code only
    when it has strictly fewer crossings than the direct after-passage code;
    otherwise the direct after-passage code is kept.  Returns a list of dicts and
    a parallel list of chosen DT codes (plain lists) for the second pass.

    ``backtrack_rounds`` > 0 (V3.4) enables backtrack-assisted simplification so
    each SnapPy-simplified crossing count escapes ``simplify('global')`` plateaus.

    ``flips`` (V4.3): the flip vector of ``dt_code`` (spherogram's 'over'
    convention).  Then every link is built from (DT, flips), every SnapPy
    export is a verified (DT, flips) pair, a passage keeps the vector, the DT
    columns carry their vectors, three ``dt_flip_vector_*`` columns are added,
    and the return value is ``(results, chosen_codes, chosen_flips)``.
    Without it: exactly as before, ``(results, chosen_codes)``.

    V4.3 (2026-10-19, round 3) -- a BARE code (``flips`` None) is run by
    ``_strand_passage_carried``: ``carry`` is the realisation the run carries
    for ``dt_code`` (spherogram's reading of a bare root, or what an earlier
    pass handed on), ``carry_active`` whether it is in force (a code the
    composite gate flags was met).  Codes the gate flags keep the
    realisation and are written with their vectors; every other code, and
    every row of a run in which nothing is flagged, exactly as before.
    ``return_meta`` adds a third return value, one dict per row for the
    chosen code: ``flips`` (the vector shown, or None), ``carry``, ``active``.
    """
    if flips is None:
        return _strand_passage_carried(
            snappy, dt_code, orig_components=orig_components,
            orig_crossings=orig_crossings, backtrack_rounds=backtrack_rounds,
            backtrack_steps=backtrack_steps, carry=carry,
            carry_active=carry_active, return_meta=return_meta)
    results = []
    chosen_codes = []
    chosen_flips = []
    with_flips = flips is not None
    if with_flips:
        flips = [bool(f) for f in flips]

    # First make the source diagram agree with the baseline crossing count used
    # for this pass.  Otherwise a later backtrack can discover that the parent
    # DT code was still on a higher-crossing plateau, while this routine still
    # enumerates crossings on that unreduced parent.  That mismatch is exactly
    # how a second-pass row can appear to "increase" from a 10-crossing baseline
    # to a 12-crossing direct child.
    source_crossings = _crossing_count(dt_code)
    if with_flips:
        link = E.link_from_dt(snappy, dt_code, flips)
        link = E.backtrack_simplify(snappy, link, mode='global',
                                    rounds=backtrack_rounds,
                                    steps=backtrack_steps, archive_flips=True)
    else:
        link = _bare_snappy_link(snappy, _dt_code_to_str(dt_code))
        link = E.backtrack_simplify(snappy, link, mode='global',
                                    rounds=backtrack_rounds, steps=backtrack_steps)
    simplified_source_crossings = len(link.crossings)
    if simplified_source_crossings < source_crossings:
        try:
            if with_flips:
                simplified_source_code, simplified_source_flips = (
                    _link_code_flips(link))
            else:
                simplified_source_code = _normalize_dt_code(link.DT_code())
            if _crossing_count(simplified_source_code) == simplified_source_crossings:
                dt_code = simplified_source_code
                source_crossings = simplified_source_crossings
                if with_flips:
                    flips = simplified_source_flips
        except Exception:  # noqa: BLE001
            pass
    if orig_components is None:
        orig_components = len(link.link_components)
    if orig_crossings is None:
        orig_crossings = source_crossings

    intervals = _build_label_intervals(dt_code)

    for comp_idx, comp in enumerate(dt_code):
        for pos_idx, label in enumerate(comp):
            label_abs = abs(label)
            partners = []
            for i, (start, end) in enumerate(intervals):
                if start <= label_abs <= end and i != comp_idx:
                    partners.append(i)
            partner = partners[0] if partners else comp_idx

            # Flip the sign at this crossing (the strand passage).  A
            # crossing change does not touch the shadow, so a flip vector is
            # kept as it is (V4.3).
            new_dt = copy.deepcopy(dt_code)
            new_comp = list(comp)
            new_comp[pos_idx] = -label
            new_dt[comp_idx] = tuple(new_comp)
            direct_str = _dt_code_to_str(new_dt)
            if with_flips:
                direct_str = E.dt_flips_to_string(new_dt, flips)
            direct_crossings = _crossing_count(new_dt)

            chosen_code = new_dt  # default: keep direct after-passage DT
            chosen_flip_vec = flips
            simplified_flips = None
            direct_components = None
            visible_components = None
            topological_components = None
            hidden_split_unknots = 0
            topological_jones = "None"
            try:
                if with_flips:
                    new_link = E.link_from_dt(snappy, new_dt, flips)
                else:
                    new_link = _bare_snappy_link(snappy, direct_str)
                direct_components = len(new_link.link_components)
                if with_flips:
                    new_link = E.backtrack_simplify(
                        snappy, new_link, mode='global',
                        rounds=backtrack_rounds, steps=backtrack_steps,
                        archive_flips=True)
                else:
                    new_link = E.backtrack_simplify(
                        snappy, new_link, mode='global',
                        rounds=backtrack_rounds, steps=backtrack_steps)
                new_components = len(new_link.link_components)
                visible_components = new_components
                topological_components = direct_components
                if direct_components is not None:
                    hidden_split_unknots = max(
                        0, int(direct_components) - int(visible_components))
                snappy_crossings = len(new_link.crossings)
                if with_flips:
                    simplified_code, simplified_flips = _link_code_flips(new_link)
                    simplified_str = _with_flips(str(simplified_code),
                                                 simplified_flips)
                else:
                    simplified_code = _normalize_dt_code(new_link.DT_code())
                    simplified_str = str(simplified_code)
                try:
                    topological_jones = str(new_link.jones_polynomial())
                except Exception:  # noqa: BLE001
                    topological_jones = "None"

                # V3.2 choice rule: SnapPy only if STRICTLY fewer crossings.
                if snappy_crossings < direct_crossings:
                    chosen_source = "snappy-simplified"
                    chosen_str = simplified_str
                    chosen_crossings = snappy_crossings
                    chosen_code = simplified_code
                    chosen_flip_vec = simplified_flips
                else:
                    chosen_source = "direct-after-passage"
                    chosen_str = _with_flips(str(new_dt), flips)
                    chosen_crossings = direct_crossings
                    chosen_code = new_dt
                    chosen_flip_vec = flips
                if with_flips:
                    chosen_jones = _jones_for_dt(snappy, chosen_code,
                                                 chosen_flip_vec)
                else:
                    chosen_jones = _jones_for_dt(snappy, chosen_code)

                # Outcome category (kept from the original script; compares the
                # SnapPy-simplified result against the original simplified link).
                cat_parts = []
                if hidden_split_unknots:
                    noun = "unknot" if hidden_split_unknots == 1 else "unknots"
                    cat_parts.append(
                        "visible DT omits %s split %s (%s -> %s visible components)"
                        % (hidden_split_unknots, noun,
                           direct_components, visible_components))
                elif new_components < orig_components:
                    cat_parts.append('fewer components (%s -> %s)'
                                     % (orig_components, new_components))
                if chosen_crossings < orig_crossings:
                    cat_parts.append('fewer crossings (%s -> %s)'
                                     % (orig_crossings, chosen_crossings))
                if (not cat_parts and new_components == orig_components
                        and chosen_crossings == orig_crossings):
                    cat_parts.append('no change')
                if (not cat_parts and (new_components > orig_components
                                       or chosen_crossings > orig_crossings)):
                    cat_parts.append('increase')
                category = ' and '.join(cat_parts)
            except Exception as exc:  # noqa: BLE001
                new_components = None
                snappy_crossings = None
                simplified_str = None
                simplified_flips = None
                chosen_source = "direct-after-passage"
                chosen_str = _with_flips(str(new_dt), flips)
                chosen_crossings = direct_crossings
                chosen_code = new_dt
                chosen_flip_vec = flips
                if with_flips:
                    chosen_jones = _jones_for_dt(snappy, chosen_code, flips)
                else:
                    chosen_jones = _jones_for_dt(snappy, chosen_code)
                topological_jones = str(chosen_jones)
                category = 'error: %s' % exc

            row = {
                'flipped_crossing': label,
                'crossing_components': (comp_idx, partner),
                'DT_code_before_simplify': direct_str,
                'DT_code_after_simplify': simplified_str,
                'DT_code_chosen': chosen_str,
                'chosen_source': chosen_source,
                'direct_crossings': direct_crossings,
                'snappy_crossings': snappy_crossings,
                'chosen_crossings': chosen_crossings,
                'Jones_polynomial': str(chosen_jones),
                'topological_Jones_polynomial': str(topological_jones),
                'orig_components': orig_components,
                'direct_components': direct_components,
                'new_components': new_components,
                'visible_components': visible_components,
                'topological_components': topological_components,
                'hidden_split_unknots': hidden_split_unknots,
                'orig_crossings': orig_crossings,
                'new_crossings': snappy_crossings,
                'outcome': category,
            }
            if with_flips:
                # V4.3: the realisations, next to the DT columns.  Named apart
                # from 'flipped_crossing', which is the CROSSING CHANGE.
                vec_cols = {
                    'dt_flip_vector_before_simplify': _flip_vector_text(flips),
                    'dt_flip_vector_after_simplify':
                        _flip_vector_text(simplified_flips),
                    'dt_flip_vector_chosen': _flip_vector_text(chosen_flip_vec),
                }
                ordered = {}
                for key, value in row.items():
                    ordered[key] = value
                    if key == 'DT_code_chosen':
                        ordered.update(vec_cols)
                row = ordered
            results.append(row)
            chosen_codes.append(chosen_code)
            chosen_flips.append(chosen_flip_vec)
    if with_flips:
        return results, chosen_codes, chosen_flips
    return results, chosen_codes


_VEC_COLUMNS = ('dt_flip_vector_before_simplify',
                'dt_flip_vector_after_simplify', 'dt_flip_vector_chosen')


def _gate_flags(code):
    """V4.3 (round 3): True when the composite gate flags the bare code."""
    try:
        return bool(E.composite_danger(code)[0])
    except Exception:  # noqa: BLE001
        return False


def _insert_vec_columns(row, values):
    """``row`` with the three ``dt_flip_vector_*`` columns right after
    'DT_code_chosen' (where a flips run puts them)."""
    ordered = {}
    for key, value in row.items():
        ordered[key] = value
        if key == 'DT_code_chosen':
            ordered.update(zip(_VEC_COLUMNS, values))
    return ordered


def _strand_passage_carried(snappy, dt_code, orig_components=None,
                            orig_crossings=None, backtrack_rounds=0,
                            backtrack_steps=20, carry=None, carry_active=False,
                            return_meta=False):
    """``strand_passage_nongui`` for a BARE code (V4.3, 2026-10-19, round 3).

    A bare run re-reads codes, and a bare code fixes a composite diagram only
    up to reflecting its parts independently: re-read, it can name another
    link (the granny for the square knot), and its passages then read as the
    wrong link's.  So the run CARRIES a realisation (``carry``: spherogram's
    reading of the root, or the vector an earlier pass handed on) and keeps
    it wherever the composite gate flags a code -- the code is then written
    with its vector, in the flips spelling, with the three
    ``dt_flip_vector_*`` columns ('' on the row's other cells) -- and,
    from there on, everywhere below it (``carry_active``: a passage, a
    SnapPy simplification and the next pass all follow the realisation;
    the codes the gate does not flag are still written bare).  This covers
    every re-read of the old bare branch: the source simplification (a root
    re-rooted onto a flagged code follows the simplified link's own vector,
    and the whole pass then follows it), each passage (a flagged direct code
    follows the carried vector, not spherogram's reading of it), SnapPy's
    simplified code (flagged: the simplified link's own vector) and the
    chosen code and its Jones polynomial.

    While nothing is flagged -- every run on a prime link with no composite
    passage -- every step is exactly the old bare branch's (the backtrack's
    ``archive_flips="auto"`` equals the default then), so the results are
    unchanged.  Returns ``(results, chosen_codes)``, plus ``meta`` (one dict
    per row: ``flips`` shown, ``carry``, ``active``) when ``return_meta``."""
    results = []
    chosen_codes = []
    meta = []
    carry = list(carry) if carry is not None else None
    active = bool(carry_active and carry is not None)
    any_shown = False

    source_crossings = _crossing_count(dt_code)
    if active:
        link = E.link_from_dt(snappy, dt_code, carry)
        link = E.backtrack_simplify(snappy, link, mode='global',
                                    rounds=backtrack_rounds,
                                    steps=backtrack_steps, archive_flips=True)
    else:
        link = _bare_snappy_link(snappy, _dt_code_to_str(dt_code))
        link = E.backtrack_simplify(snappy, link, mode='global',
                                    rounds=backtrack_rounds,
                                    steps=backtrack_steps, archive_flips="auto")
    simplified_source_crossings = len(link.crossings)
    if simplified_source_crossings < source_crossings:
        try:
            if active:
                simplified_source_code, simplified_source_carry = (
                    _link_code_flips(link))
                rerooted_active = True
            else:
                simplified_source_code = _normalize_dt_code(link.DT_code())
                simplified_source_carry = None
                rerooted_active = False
                if _gate_flags(simplified_source_code):
                    # the source simplified to a flagged code: re-root on it
                    # WITH the simplified link's own vector, and follow it
                    try:
                        simplified_source_code, simplified_source_carry = (
                            _link_code_flips(link))
                        rerooted_active = True
                    except ValueError:
                        simplified_source_carry = None
                if not rerooted_active:
                    simplified_source_carry = E._own_vector(
                        link, simplified_source_code)
            if _crossing_count(simplified_source_code) == simplified_source_crossings:
                if rerooted_active and not active:
                    print("[info] the source simplified to a code the "
                          "composite gate flags; the pass follows the "
                          "simplified link's own realisation, %s"
                          % E.dt_flips_to_string(simplified_source_code,
                                                 simplified_source_carry))
                dt_code = simplified_source_code
                source_crossings = simplified_source_crossings
                carry = simplified_source_carry
                active = bool(rerooted_active and carry is not None)
        except Exception:  # noqa: BLE001
            pass
    if orig_components is None:
        orig_components = len(link.link_components)
    if orig_crossings is None:
        orig_crossings = source_crossings

    intervals = _build_label_intervals(dt_code)

    for comp_idx, comp in enumerate(dt_code):
        for pos_idx, label in enumerate(comp):
            label_abs = abs(label)
            partners = []
            for i, (start, end) in enumerate(intervals):
                if start <= label_abs <= end and i != comp_idx:
                    partners.append(i)
            partner = partners[0] if partners else comp_idx

            # Flip the sign at this crossing (the strand passage).  A
            # crossing change does not touch the shadow, so the carried
            # vector is kept as it is.
            new_dt = copy.deepcopy(dt_code)
            new_comp = list(comp)
            new_comp[pos_idx] = -label
            new_dt[comp_idx] = tuple(new_comp)
            direct_bare = _dt_code_to_str(new_dt)
            direct_crossings = _crossing_count(new_dt)
            direct_flagged = _gate_flags(new_dt)
            dvec = None                      # the realisation followed here
            if active or direct_flagged:
                dvec = carry
                if dvec is None and direct_flagged:
                    dvec = E.spherogram_reading(new_dt)[0]
                if dvec is not None and len(dvec) == 0:
                    dvec = None
            row_active = dvec is not None
            direct_shown = direct_flagged and dvec is not None
            direct_str = (E.dt_flips_to_string(new_dt, dvec) if direct_shown
                          else direct_bare)

            chosen_code = new_dt  # default: keep direct after-passage DT
            chosen_vec = dvec if dvec is not None else carry
            chosen_shown = direct_shown
            chosen_active = row_active
            simplified_flips = None
            simplified_shown = False
            direct_components = None
            visible_components = None
            topological_components = None
            hidden_split_unknots = 0
            topological_jones = "None"
            try:
                if dvec is not None:
                    new_link = E.link_from_dt(snappy, new_dt, dvec)
                else:
                    new_link = _bare_snappy_link(snappy, direct_bare)
                direct_components = len(new_link.link_components)
                new_link = E.backtrack_simplify(
                    snappy, new_link, mode='global',
                    rounds=backtrack_rounds, steps=backtrack_steps,
                    archive_flips=(True if dvec is not None else "auto"))
                new_components = len(new_link.link_components)
                visible_components = new_components
                topological_components = direct_components
                if direct_components is not None:
                    hidden_split_unknots = max(
                        0, int(direct_components) - int(visible_components))
                snappy_crossings = len(new_link.crossings)
                if dvec is not None:
                    simplified_code, simplified_flips = _link_code_flips(new_link)
                    simplified_shown = _gate_flags(simplified_code)
                    simplified_active = True
                else:
                    simplified_code = _normalize_dt_code(new_link.DT_code())
                    simplified_active = False
                    if _gate_flags(simplified_code):
                        # SnapPy's simplified code is flagged: written with
                        # the simplified link's own vector, never re-read
                        try:
                            simplified_code, simplified_flips = (
                                _link_code_flips(new_link))
                            simplified_shown = simplified_active = True
                        except ValueError:
                            simplified_flips = None
                    if not simplified_active:
                        simplified_flips = E._own_vector(new_link,
                                                         simplified_code)
                if simplified_shown:
                    simplified_str = _with_flips(str(simplified_code),
                                                 simplified_flips)
                    if simplified_str == str(simplified_code):
                        simplified_shown = False     # no crossings
                else:
                    simplified_str = str(simplified_code)
                try:
                    topological_jones = str(new_link.jones_polynomial())
                except Exception:  # noqa: BLE001
                    topological_jones = "None"

                # V3.2 choice rule: SnapPy only if STRICTLY fewer crossings.
                if snappy_crossings < direct_crossings:
                    chosen_source = "snappy-simplified"
                    chosen_str = simplified_str
                    chosen_crossings = snappy_crossings
                    chosen_code = simplified_code
                    chosen_vec = simplified_flips
                    chosen_shown = simplified_shown
                    chosen_active = simplified_active
                else:
                    chosen_source = "direct-after-passage"
                    chosen_str = (_with_flips(str(new_dt), dvec)
                                  if direct_shown else str(new_dt))
                    chosen_crossings = direct_crossings
                    chosen_code = new_dt
                    chosen_vec = dvec if dvec is not None else carry
                    chosen_shown = direct_shown
                    chosen_active = row_active
                if chosen_active and chosen_vec is not None:
                    chosen_jones = _jones_for_dt(snappy, chosen_code, chosen_vec)
                else:
                    chosen_jones = _jones_for_dt(snappy, chosen_code)

                # Outcome category (kept from the original script; compares the
                # SnapPy-simplified result against the original simplified link).
                cat_parts = []
                if hidden_split_unknots:
                    noun = "unknot" if hidden_split_unknots == 1 else "unknots"
                    cat_parts.append(
                        "visible DT omits %s split %s (%s -> %s visible components)"
                        % (hidden_split_unknots, noun,
                           direct_components, visible_components))
                elif new_components < orig_components:
                    cat_parts.append('fewer components (%s -> %s)'
                                     % (orig_components, new_components))
                if chosen_crossings < orig_crossings:
                    cat_parts.append('fewer crossings (%s -> %s)'
                                     % (orig_crossings, chosen_crossings))
                if (not cat_parts and new_components == orig_components
                        and chosen_crossings == orig_crossings):
                    cat_parts.append('no change')
                if (not cat_parts and (new_components > orig_components
                                       or chosen_crossings > orig_crossings)):
                    cat_parts.append('increase')
                category = ' and '.join(cat_parts)
            except Exception as exc:  # noqa: BLE001
                new_components = None
                snappy_crossings = None
                simplified_str = None
                simplified_flips = None
                simplified_shown = False
                chosen_source = "direct-after-passage"
                chosen_str = (_with_flips(str(new_dt), dvec)
                              if direct_shown else str(new_dt))
                chosen_crossings = direct_crossings
                chosen_code = new_dt
                chosen_vec = dvec if dvec is not None else carry
                chosen_shown = direct_shown
                chosen_active = row_active
                if dvec is not None:
                    chosen_jones = _jones_for_dt(snappy, chosen_code, dvec)
                else:
                    chosen_jones = _jones_for_dt(snappy, chosen_code)
                topological_jones = str(chosen_jones)
                category = 'error: %s' % exc

            row = {
                'flipped_crossing': label,
                'crossing_components': (comp_idx, partner),
                'DT_code_before_simplify': direct_str,
                'DT_code_after_simplify': simplified_str,
                'DT_code_chosen': chosen_str,
                'chosen_source': chosen_source,
                'direct_crossings': direct_crossings,
                'snappy_crossings': snappy_crossings,
                'chosen_crossings': chosen_crossings,
                'Jones_polynomial': str(chosen_jones),
                'topological_Jones_polynomial': str(topological_jones),
                'orig_components': orig_components,
                'direct_components': direct_components,
                'new_components': new_components,
                'visible_components': visible_components,
                'topological_components': topological_components,
                'hidden_split_unknots': hidden_split_unknots,
                'orig_crossings': orig_crossings,
                'new_crossings': snappy_crossings,
                'outcome': category,
            }
            if direct_shown or simplified_shown or chosen_shown:
                # the realisations of the codes the gate flags, next to the
                # DT columns (named apart from 'flipped_crossing', which is
                # the CROSSING CHANGE)
                any_shown = True
                row = _insert_vec_columns(row, (
                    _flip_vector_text(dvec) if direct_shown else "",
                    _flip_vector_text(simplified_flips) if simplified_shown else "",
                    _flip_vector_text(chosen_vec) if chosen_shown else ""))
            results.append(row)
            chosen_codes.append(chosen_code)
            meta.append({
                "flips": list(chosen_vec) if (chosen_shown and chosen_vec) else None,
                "carry": list(chosen_vec) if chosen_vec is not None else None,
                "active": bool(chosen_active and chosen_vec is not None),
            })
    if any_shown:
        # one set of columns per sheet: '' where a row's codes are unflagged
        results = [row if _VEC_COLUMNS[0] in row
                   else _insert_vec_columns(row, ("", "", ""))
                   for row in results]
    if return_meta:
        return results, chosen_codes, meta
    return results, chosen_codes


# ---- structure fingerprinting + merged passage tree (for the overview SVG) --
def _num_components(snappy, dt_code, flips=None):
    """Number of link components of a DT code (as drawn, no simplification).

    V4.3: with ``flips`` the link is built from (DT, flips)."""
    try:
        if flips is not None:
            return len(E.link_from_dt(snappy, dt_code, flips).link_components)
        return len(_bare_snappy_link(snappy, _dt_code_to_str(dt_code)).link_components)
    except Exception:  # noqa: BLE001
        return -1


def _jones_for_dt(snappy, dt_code, flips=None):
    """Jones polynomial for the exact DT code shown in the spreadsheet.

    V4.3: with ``flips`` it is the Jones polynomial of THAT realisation (a
    bare code lets spherogram choose one, which for a composite can be a
    different link: the square knot's code read bare is the granny)."""
    try:
        if flips is not None:
            return str(E.link_from_dt(snappy, dt_code, flips).jones_polynomial())
        return str(_bare_snappy_link(snappy, _dt_code_to_str(dt_code)).jones_polynomial())
    except Exception:  # noqa: BLE001
        return "None"


def _flip_vector_text(flips):
    """'[0,1,0,0,0,1]' (the flips spelling's vector part); '' for None."""
    if flips is None:
        return ""
    return "[%s]" % ",".join("1" if f else "0" for f in flips)


def _with_flips(text, flips):
    """A DT text with its vector appended in the toolkit's spelling (bare text
    for None, and for an empty vector: no crossings, nothing to realise)."""
    if text is None or flips is None or len(flips) == 0:
        return text
    return format_dt_flips(text, flips)


def _text_carries_vector(text):
    """True when DT text carries a flip vector -- or text after the code that
    is not one (which ``split_dt_flips`` refuses): errors on such input are
    reported as one [error] line, while a bare code fails exactly as before."""
    try:
        return E._carries_flips(text)
    except ValueError:
        return True


def _unparseable_text(dt_text, exc):
    """The one-line refusal of a BARE DT code ``E.parse_dt_flips`` cannot read
    (round 4d, 2026-10-20: '--nongui' and '--demo' ended in a traceback on an
    alphabetical 'DT[daddcab]' or a prefix-less '8,6,2,4').  The scorer's
    wording (score_diagram's _unparseable_error): what this tool reads, and
    that dt_converter.py converts other spellings.  Object addresses in the
    parser's message ('<ast.Name object at 0x...>') are dropped, so the line
    is the same in every process; a '# comment' is not quoted."""
    why = " ".join(re.sub(r" at 0x[0-9A-Fa-f]+", "", str(exc)).split()) or type(exc).__name__
    try:
        text = split_dt_flips(str(dt_text))[0].strip()
    except ValueError:
        text = str(dt_text).strip()
    return ("strand passage cannot parse the DT code %s (%s): it reads signed DT codes written "
            "as lists of nonzero even integers, e.g. 'DT: [(4, 6, 2)]', optionally followed by a "
            "flip vector.  dt_converter.py -i \"%s\" --to dt converts other spellings "
            "(alphabetical, packed hex, Gauss, PD, braid words, names) to that one (--to "
            "dt-flips adds the flip vector) and says what is wrong with a code it cannot read"
            % (text, why[:160], text))


def _link_code_flips(link):
    """(code, flips) of a link in ``_normalize_dt_code``'s shape ([] for no
    crossings), with a verified vector (``E.link_to_dt_flips``)."""
    comps, flips = E.link_to_dt_flips(link)
    return [tuple(int(x) for x in comp) for comp in comps], list(flips)


def _bare_root_realisation(comps, negative_even="over", follow_unreadable=False):
    """V4.3 (2026-10-19, round 3): the realisation a BARE root is followed
    with, ``(vector, flagged, line)``.

    A bare signed DT code fixes a composite diagram only up to reflecting its
    parts independently, and spherogram resolves it by its own rule, so a
    run that re-reads bare codes can drift to another link (the granny's
    passages read as the square knot's; a prime-looking diagram of the
    square knot re-rooted onto a composite code read as the granny).  Every
    bare root therefore CARRIES a realisation through the run -- spherogram's
    own reading of it (``E.root_realisation``: what ``snappy.Link(bare)``
    builds, i.e. what a bare run analyses first), or dt_converter's mended
    reading for a code the composite gate flags whose reading spherogram
    gets illegal (split codes) -- and the run writes it out only on codes
    the gate flags.

    * ``vector`` -- the carried vector, or None (spherogram cannot read the
      code: the run reads it bare, as before);
    * ``flagged`` -- the composite gate flags the root;
    * ``line`` -- None for an unflagged root (it runs exactly as before);
      '[info] ...' for a flagged root followed with ``vector``; '[warn] ...'
      for a flagged root with no reading (read bare as before).

    ``comps`` is read with ``negative_even``; a vector describes the shadow,
    so it is the same in both conventions.

    Round 4 (2026-10-19): a flagged root spherogram's decoder cannot read at
    all is followed too (``E.root_realisation`` mends that stage now, as the
    scorer runs such roots): with dt_converter's reading (the smallest flip
    vector spherogram builds) or a legal reading spherogram builds
    (``E.buildable_legal_flips``).  ``follow_unreadable`` (--nongui, which
    has no bare fallback) does the same for an UNFLAGGED root, with an
    '[info]' line, and returns an '[error]' line instead of the '[warn]'
    where a root has no reading at all."""
    if negative_even == "under":
        over = [tuple(-int(x) for x in comp) for comp in comps]
    else:
        over = [tuple(int(x) for x in comp) for comp in comps]
    vector, flagged, why, how = E.root_realisation(
        over, mend_unflagged=bool(follow_unreadable))
    if vector is not None:
        vector = list(vector)
    if not flagged:
        if not follow_unreadable or (vector is not None
                                     and how == "spherogram's reading"):
            return vector, False, None
        if vector is None:
            nowhere = E.no_diagram_reason(over)
            if nowhere:
                return None, False, (
                    "[error] %s: %s, so --nongui has no link to pass strands "
                    "through; check the code" % (_dt_code_to_str(over), nowhere))
            empty = [i + 1 for i, c in enumerate(over) if not len(c)]
            if empty and len(empty) < len(over):
                # round 4d (2026-10-20): the cause is the empty component
                # (dt_converter refuses it, and spherogram builds no
                # realisation with a TRAILING one: IndexError), so a flip
                # vector is not the remedy; without the component spherogram
                # reads the code as it read this one, and --nongui follows it
                return None, False, (
                    "[error] %s: %s, so --nongui has no link to pass strands "
                    "through; give the code without its empty component%s %s "
                    "(spherogram's DT decoder skips empty components)"
                    % (_dt_code_to_str(over), how, "s" if len(empty) > 1 else "",
                       ", ".join(str(i) for i in empty)))
            return None, False, (
                "[error] %s: %s, so --nongui has no link to pass strands "
                "through; give the code with a flip vector"
                % (_dt_code_to_str(over), how))
        return vector, False, (
            "[info] spherogram cannot read the bare code %s, so the run "
            "follows %s: %s, through every step (the composite gate does not "
            "flag it, so every realisation of it is the same link up to mirror "
            "image; the root is written with this vector, every other code the "
            "gate does not flag bare)"
            % (_dt_code_to_str(over), _reading_phrase(how),
               E.dt_flips_to_string(over, vector)))
    if not vector:
        hint = E.legal_flips_hint(over)
        return None, True, ("[%s] the composite gate flags the bare code (%s), "
                            "but %s, so %s; name its realisation with a flip "
                            "vector%s"
                            % ("error" if follow_unreadable else "warn", why, how,
                               "--nongui has no link to pass strands through"
                               if follow_unreadable else "it is read bare as before",
                               (", e.g. %s" % hint) if hint else ""))
    if how == "spherogram's reading":
        which = "spherogram's reading of it"
    elif how.startswith("dt_converter's mended reading"):
        which = ("dt_converter's mended reading of it (spherogram's own "
                 "reading describes no diagram on the sphere; the pieces it "
                 "read wrong take a legal vector, first crossing positive)")
    else:
        which = _reading_phrase(how)       # round 4: a decoder-stage reading
    return vector, True, (
        "[info] the composite gate flags the bare code (%s); following %s, "
        "%s%s, through every step (codes the gate flags are written with "
        "their flip vectors; give the code with a flip vector to follow "
        "another realisation)"
        % (why, which, E.dt_flips_to_string(over, vector),
           " in spherogram's 'over' convention"
           if negative_even == "under" else ""))


def _reading_phrase(how):
    """Round 4 (2026-10-19): ``E.root_realisation``'s ``how`` for a root
    spherogram's decoder cannot read, as the [info] lines name it: dt_converter's
    reading as written, a reading from ``E.buildable_legal_flips`` as 'a legal
    reading of it (...)', never as dt_converter's."""
    how = str(how)
    if how.startswith("a legal flip vector"):
        desc, sep, rest = how.partition(", since ")
        return "a legal reading of it (%s)%s%s" % (desc, sep, rest)
    return how


def _refuse_unrealisable(comps):
    """Round 4 (2026-10-19): raise a user-facing ValueError for a BARE code no
    flip vector makes a diagram on the sphere (``E.no_diagram_reason``:
    deterministic, no decode) before a run builds or draws it; any other
    code passes.  It used to end in a traceback ('Link isn't planar', 'The
    crossing graph is not planar'); the GUI shows it in its error window and
    ``--demo`` prints it as one line."""
    reason = getattr(E, "no_diagram_reason", None)
    why = reason(comps) if callable(reason) else None
    if why:
        err = ValueError("%s: %s, so there is no diagram to draw; check the code"
                         % (_dt_code_to_str(comps), why))
        err.user_facing = True
        raise err


def _bare_snappy_link(snappy, text):
    """``snappy.Link(text)`` for a BARE code, after ``E._guard_dt_decoder()``
    (round 4b, 2026-10-20): dt_converter's guard on spherogram's decoder is
    installed when the engine is imported, and retried here if that failed
    (under sage with faulthandler on, the first import of spherogram raises);
    a no-op once installed, so the link is exactly ``snappy.Link(text)``."""
    E._guard_dt_decoder()
    return snappy.Link(text)


def _bare_root_unbuildable(snappy, dt_code):
    """Round 4 (2026-10-19): '' when ``snappy.Link`` builds the bare root, else
    why not ('<ExceptionType>: <message>'; dt_converter's guard makes the
    decoder raise DecoderLoop where it would loop) -- for --nongui, which used
    to end in a traceback (or never return) on such a root."""
    try:
        _bare_snappy_link(snappy, _dt_code_to_str(dt_code))
    except Exception as exc:  # noqa: BLE001
        return "%s: %s" % (type(exc).__name__, str(exc)[:160])
    return ""


def _bare_root_flips(comps, negative_even="over"):
    """V4.3 (2026-10-19): ``(vector, line)`` for a bare root the composite
    gate flags (``_bare_root_realisation``'s vector and line), ``(None,
    None)`` for one it does not flag; ``(None, '[warn] ...')`` when a flagged
    root has no reading."""
    vector, flagged, line = _bare_root_realisation(comps, negative_even)
    if not flagged:
        return None, None
    return vector, line


def _label_sort_key(s):
    t = str(s).lstrip("-")
    return int(s) if t.isdigit() else 0


def _safe_sheet_name(base, used):
    """Excel-safe unique sheet name, limited to 31 characters."""
    cleaned = "".join("_" if ch in r'[]:*?/\\' else ch for ch in str(base))
    cleaned = cleaned.strip() or "sheet"
    root = cleaned[:31]
    name = root
    i = 2
    while name in used:
        suffix = "_%d" % i
        name = root[:31 - len(suffix)] + suffix
        i += 1
    used.add(name)
    return name


def _module_version(module):
    """Best-effort version string for a runtime dependency."""
    for attr in ("VERSION", "__version__", "version"):
        try:
            value = getattr(module, attr)
            if callable(value):
                value = value()
            if value is not None:
                return str(value)
        except Exception:  # noqa: BLE001
            pass
    return "unknown"


def _runtime_command():
    """Reconstruct the command that is useful to paste into a terminal."""
    prefix = ["sage", "-python"] if _running_under_sage() else [sys.executable]
    return shlex.join(prefix + list(sys.argv))


def _run_info_rows(snappy, pandas_module, dt_string, dt_code, out_path,
                   overview_path, negative_even, backtrack_rounds,
                   backtrack_steps, crossing_order, crossing_map,
                   crossing_input, drawing_options, drawing_session_path,
                   first_pass_rows, continuable_first_passages,
                   merged_first_step_continuations, second_pass_sheets,
                   second_pass_rows, flips=None, flips_note=None,
                   negative_even_note=None):
    """Workbook metadata for reproducible nongui output.

    V4.3: with a flip vector, ``normalized_input_dt`` is in the flips spelling
    and rows ``input_dt_flip_vector`` (+ ``dt_flip_vector_note``) follow it.
    A bare code rewritten from 'under' into the 'over' convention gets a row
    ``negative_even_note`` after ``negative_even`` (``normalized_input_dt`` is
    then the rewritten code)."""
    rows = [
        ("created_at", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("software", "DT Link Toolkit"),
        ("software_version", VERSION),
        ("entry_script", os.path.basename(sys.argv[0]) if sys.argv else ""),
        ("engine_module", getattr(E, "__name__", "unknown")),
        ("engine_version", getattr(E, "VERSION", "unknown")),
        ("drawing_module", getattr(D, "__name__", "unknown")),
        ("drawing_version", getattr(D, "VERSION", "unknown")),
        ("python_version", sys.version.replace("\n", " ")),
        ("python_executable", sys.executable),
        ("running_under_sage", str(_running_under_sage())),
        ("snappy_version", _module_version(snappy)),
        ("pandas_version", _module_version(pandas_module)),
        ("command", _runtime_command()),
        ("argv", repr(sys.argv)),
        ("working_directory", os.getcwd()),
        ("input_dt", dt_string),
        ("normalized_input_dt", _dt_code_to_str(dt_code)),
        ("negative_even", negative_even),
    ]
    if negative_even_note is not None:
        rows.append(("negative_even_note", negative_even_note))
    rows += [
        ("output_xlsx", out_path),
        ("overview_svg", overview_path),
        ("backtrack_rounds", int(backtrack_rounds or 0)),
        ("backtrack_steps", int(backtrack_steps or 0)),
        ("drawing_session", drawing_session_path or ""),
        ("drawing_options", repr(normalize_drawing_options(drawing_options))),
        ("crossing_input", crossing_input or ""),
        ("crossing_order", crossing_order or ""),
        ("crossing_map", crossing_map or ""),
        ("second_pass_continuation_criterion", NONGUI_SECOND_PASS_CRITERION),
        ("first_pass_rows", int(first_pass_rows)),
        ("continuable_first_step_passages", int(continuable_first_passages)),
        ("merged_first_step_continuations", int(merged_first_step_continuations)),
        ("second_pass_sheets", int(second_pass_sheets)),
        ("second_pass_rows", int(second_pass_rows)),
    ]
    if flips is not None:
        out = []
        for key, value in rows:
            if key == "normalized_input_dt":
                value = E.dt_flips_to_string(dt_code, flips)
            out.append((key, value))
            if key == "normalized_input_dt":
                out.append(("input_dt_flip_vector", _flip_vector_text(flips)))
                out.append(("dt_flip_vector_note", flips_note or (
                    "bit k <-> the k-th DT entry (odd label 2k+1); walking "
                    "along the visit with the smaller label, 0 = the other "
                    "strand crosses from right to left (spherogram/SnapPy "
                    "convention); every link here is built from (DT, flips)")))
        rows = out
    return rows


def _edge_list_from_map(edge_map):
    """Convert the overview edge accumulator into a sorted serializable list."""
    return [{"src": s, "dst": d, "labels": sorted(l, key=_label_sort_key)}
            for (s, d), l in edge_map.items()]


# --------------------------------------------------------------------------- #
#  V4.2: make the exported SVG's font declarations portable
# --------------------------------------------------------------------------- #
# With ``svg.fonttype: none`` Matplotlib keeps text as real SVG <text> (which is
# what makes the overview editable in Illustrator), but it declares the font with
# the CSS *shorthand*:
#
#     style="fill: #111827; font: 700 8.6px 'Arial'"
#
# Two problems follow.  The shorthand resets every font property it does not
# name to its initial value, and a consumer that implements it only partially
# can lose the size or the weight entirely -- Illustrator is the one that bites
# in practice.  And 'Arial' is named with no generic fallback, so a machine
# without Arial substitutes silently.  Either way the consumer re-lays-out the
# string; since each run carries one anchor and each background box is a separate
# path sized from Matplotlib's measurements, the text then slides out of its box.
#
# Rewriting the shorthand as longhand fixes both without touching geometry, text
# or anchors.
_SVG_FONT_SHORTHAND = re.compile(
    r"font:\s*"
    r"(?:(?P<style>italic|oblique)\s+)?"
    r"(?:(?P<weight>\d{3}|bold|normal|lighter|bolder)\s+)?"
    r"(?P<size>[\d.]+)px\s+"
    r"'(?P<family>[^']+)'")

_SVG_FONT_FALLBACKS = {
    "Arial": "Arial, 'Helvetica Neue', Helvetica, sans-serif",
    "Helvetica": "Helvetica, Arial, sans-serif",
    "DejaVu Sans": "'DejaVu Sans', Verdana, sans-serif",
    "DejaVu Sans Mono": "'DejaVu Sans Mono', Menlo, Consolas, monospace",
}


def _svg_font_longhand(match):
    """Expand one CSS ``font:`` shorthand match into explicit longhand."""
    family = match.group("family")
    stack = _SVG_FONT_FALLBACKS.get(family)
    if stack is None:
        quoted = "'%s'" % family if " " in family else family
        stack = "%s, sans-serif" % quoted
    parts = ["font-family: %s" % stack, "font-size: %spx" % match.group("size")]
    if match.group("weight"):
        parts.append("font-weight: %s" % match.group("weight"))
    if match.group("style"):
        parts.append("font-style: %s" % match.group("style"))
    return "; ".join(parts)


def _rewrite_svg_font_declarations(path):
    """Replace CSS ``font:`` shorthand with longhand in an exported SVG.

    Returns the number of declarations rewritten.  Never raises: the figure is
    already on disk and correct, so a problem here must not cost the user the
    file.
    """
    try:
        with open(path, "r", encoding="utf-8") as fh:
            svg = fh.read()
        rewritten, count = _SVG_FONT_SHORTHAND.subn(_svg_font_longhand, svg)
        if count:
            with open(path, "w", encoding="utf-8") as fh:
                fh.write(rewritten)
        return count
    except Exception as exc:  # noqa: BLE001
        print("[warn] could not rewrite SVG font declarations in %s: %s"
              % (path, exc), file=sys.stderr)
        return 0


# --------------------------------------------------------------------------- #
#  V4.1: shade workbook rows by the overview-SVG card they are drawn as
# --------------------------------------------------------------------------- #
# Light, print-safe backgrounds.  They only have to be distinguishable from each
# other and leave black text readable; they are NOT the overview card colours
# (those encode the outcome vs. the original, which several different cards
# share, so they cannot identify a card).
_CARD_SHADES = (
    "FFF3C4", "D8ECF7", "E2F0D9", "F8DDE3", "E7E0F4", "FCE3CB",
    "D9F2EC", "F3E1F7", "E8EAF6", "F0EAD6", "DCE9F5", "F6E0E0",
    "E3F1D5", "EFE3D0", "DDEEF0", "F5E6F0", "E6EFDA", "F2E8DC",
)
_ROOT_CARD_SHADE = "EFF6FF"   # matches the blue depth-0 card in the SVG


def _card_shade(card_id, node_by_id=None):
    """Background colour for one overview card (stable for a given workbook)."""
    if node_by_id is not None:
        node = node_by_id.get(card_id)
        if node is not None and int(node.get("depth", 1)) == 0:
            return _ROOT_CARD_SHADE
    return _CARD_SHADES[int(card_id) % len(_CARD_SHADES)]


def _card_index_rows(node_list, edge_list, card_sheets, original_crossings):
    """One legend row per overview card, in the order the SVG lays them out."""
    arriving = {}
    for edge in edge_list:
        arriving.setdefault(edge["dst"], []).extend(str(x) for x in edge["labels"])

    by_depth = {}
    for node in node_list:
        by_depth.setdefault(int(node["depth"]), []).append(node)
    step_name = {0: "Original", 1: "After 1st strand passage",
                 2: "After 2nd strand passage"}

    rows = []
    any_flips = any(node.get("flips") is not None for node in node_list)
    for depth in sorted(by_depth):
        # Same order the cards are stacked in each SVG column.
        for node in sorted(by_depth[depth],
                           key=lambda x: (x["n_crossings"], x["id"])):
            _b, _f, tag = _outcome_style(node, original_crossings)
            labels = sorted(arriving.get(node["id"], []), key=_label_sort_key)
            row = {
                "svg_card_id": node["id"],
                "svg_column": step_name.get(depth, "step %d" % depth),
                "depth": depth,
                "DT_code_drawn": node["dt_str"],
                "crossings": node["n_crossings"],
                "components": node["n_components"],
                "Jones_polynomial": node["jones"],
                "outcome_vs_original": tag or "original",
                "passages_arriving": len(labels),
                "arriving_flips": ", ".join(labels),
                "workbook_sheets": ", ".join(
                    sorted(card_sheets.get(node["id"], set()))),
            }
            if any_flips:
                # V4.3: the card's realisation ('arriving_flips' counts
                # crossing changes, not this vector); round 3: '' on the
                # cards of a bare run whose codes the gate does not flag
                ordered = {}
                for key, value in row.items():
                    ordered[key] = value
                    if key == "DT_code_drawn":
                        ordered["dt_flip_vector"] = _flip_vector_text(
                            node.get("flips"))
                row = ordered
            rows.append(row)
    return rows


def _shade_rows_by_card(worksheet, card_ids, n_columns, node_by_id=None):
    """Fill each data row of ``worksheet`` with its card's shade.

    ``card_ids[i]`` is the card for data row ``i`` (``None`` leaves the row
    unshaded, which is what an errored passage with no drawn structure gets).
    Never raises: shading is cosmetic, so a styling problem must not cost the
    user the workbook itself.
    """
    try:
        from openpyxl.styles import PatternFill  # type: ignore
    except Exception:  # noqa: BLE001
        return 0
    fills = {}
    shaded = 0
    for offset, card_id in enumerate(card_ids):
        if card_id is None:
            continue
        colour = _card_shade(card_id, node_by_id)
        fill = fills.get(colour)
        if fill is None:
            fill = PatternFill(start_color=colour, end_color=colour,
                               fill_type="solid")
            fills[colour] = fill
        for col in range(1, int(n_columns) + 1):
            worksheet.cell(row=offset + 2, column=col).fill = fill
        shaded += 1
    return shaded


def _renumber_passage_graph(node_list, edge_list):
    """Make node ids contiguous after reconciliation drops merged-away nodes.

    Returns ``(nodes, edges, old_to_new)``.  V4.1 also returns the id mapping so
    the caller can follow a spreadsheet row to the card it ends up drawn as.
    """
    old_to_new = {}
    new_nodes = []
    for new_id, node in enumerate(node_list):
        old_id = node["id"]
        nd = dict(node)
        nd["id"] = new_id
        fp = nd.get("fp")
        if fp is not None:
            nd["fp"] = (int(nd["depth"]), int(nd["n_crossings"]),
                        int(nd["n_components"]), fp[3])
        old_to_new[old_id] = new_id
        new_nodes.append(nd)

    from collections import defaultdict
    emap = defaultdict(set)
    for edge in edge_list:
        if edge["src"] not in old_to_new or edge["dst"] not in old_to_new:
            continue
        emap[(old_to_new[edge["src"]], old_to_new[edge["dst"]])].update(
            str(x) for x in edge["labels"])
    new_edges = [{"src": s, "dst": d,
                  "labels": sorted(v, key=_label_sort_key)}
                 for (s, d), v in emap.items()]
    return new_nodes, new_edges, old_to_new


def reconcile_steps(snappy, node_list, edge_list, rounds=400, steps=30,
                    max_passes=3):
    """Targeted per-step reconciliation of incompletely-simplified structures.

    Within each step, structures that share the same Jones polynomial (up to
    mirror q->1/q and framing q^n) and component count but show DIFFERENT
    crossing counts are the same link left at different simplify('global')
    plateaus.  This drives the higher-crossing ones DOWN to the group's minimum
    with a targeted backtrack (stopping as soon as it reaches the target), which
    is far more effective than a big global round count, then re-merges the
    now-identical structures (remapping the arrows).

    Returns ``(nodes, edges, remap)``.  V4.1 also returns the id mapping (every
    input id -> the id that survived) so the caller can follow a spreadsheet row
    to the card it ends up drawn as.
    """
    from collections import defaultdict

    def gkey(nd):
        return (nd["depth"], nd["n_components"], nd["fp"][3])

    for _ in range(max_passes):
        groups = defaultdict(list)
        for nd in node_list:
            groups[gkey(nd)].append(nd)
        changed = False
        for _key, grp in groups.items():
            if len(grp) < 2:
                continue
            target = min(nd["n_crossings"] for nd in grp)
            for nd in grp:
                if nd["n_crossings"] <= target:
                    continue
                nd_flips = nd.get("flips")      # V4.3: the realisation
                carried_run = "carry" in nd     # round 3: a bare run's node
                use = None
                if carried_run and nd.get("active") and nd.get("carry"):
                    use = nd["carry"]
                try:
                    if carried_run:
                        if use is not None:
                            L = E.link_from_dt(snappy, nd["dt_code"], use)
                            L = E.backtrack_simplify(snappy, L, mode="global",
                                                     rounds=rounds, steps=steps,
                                                     target=target,
                                                     archive_flips=True)
                        else:
                            L = _bare_snappy_link(snappy, _dt_code_to_str(nd["dt_code"]))
                            L = E.backtrack_simplify(snappy, L, mode="global",
                                                     rounds=rounds, steps=steps,
                                                     target=target,
                                                     archive_flips="auto")
                    elif nd_flips is not None:
                        L = E.link_from_dt(snappy, nd["dt_code"], nd_flips)
                        L = E.backtrack_simplify(snappy, L, mode="global",
                                                 rounds=rounds, steps=steps,
                                                 target=target,
                                                 archive_flips=True)
                    else:
                        L = _bare_snappy_link(snappy, _dt_code_to_str(nd["dt_code"]))
                        L = E.backtrack_simplify(snappy, L, mode="global",
                                                 rounds=rounds, steps=steps,
                                                 target=target)
                    n = len(L.crossings)
                except Exception:  # noqa: BLE001
                    continue
                if n < nd["n_crossings"]:
                    if carried_run:
                        _reconciled_carried(nd, L, n, use)
                        changed = True
                        continue
                    if nd_flips is not None:
                        try:
                            code, nd_flips = _link_code_flips(L)
                        except Exception:  # noqa: BLE001
                            code = nd["dt_code"]
                        nd["dt_code"] = code
                        nd["flips"] = nd_flips
                        nd["n_crossings"] = n
                        nd["dt_str"] = E.dt_flips_to_string(code, nd_flips)
                        changed = True
                        continue
                    try:
                        code = _normalize_dt_code(L.DT_code())
                    except Exception:  # noqa: BLE001
                        code = nd["dt_code"]
                    nd["dt_code"] = code
                    nd["n_crossings"] = n
                    nd["dt_str"] = _dt_code_to_str(code)
                    changed = True
        if not changed:
            break

    # Re-merge structures within a step that now share crossing count too.
    remap, canon, new_nodes = {}, {}, []
    for nd in node_list:
        mk = (nd["depth"], nd["n_crossings"], nd["n_components"], nd["fp"][3])
        if mk in canon:
            remap[nd["id"]] = canon[mk]["id"]
        else:
            canon[mk] = nd
            remap[nd["id"]] = nd["id"]
            new_nodes.append(nd)

    from collections import defaultdict as _dd
    emap = _dd(set)
    for e in edge_list:
        s = remap.get(e["src"], e["src"])
        d = remap.get(e["dst"], e["dst"])
        emap[(s, d)].update(e["labels"])
    new_edges = [{"src": s, "dst": d,
                  "labels": sorted(v, key=_label_sort_key)}
                 for (s, d), v in emap.items()]
    return new_nodes, new_edges, remap


def _reconciled_carried(nd, L, n, use):
    """Update a bare run's node ``nd`` after reconciliation drove it down to
    the link ``L`` with ``n`` crossings (V4.3, round 3).  ``use`` is the
    vector the reduction followed (None: read bare, as before).  The node
    keeps the realisation it carries and writes it out only when the gate
    flags its new code; an unflagged node read bare is updated exactly as
    before (and carries ``L``'s own vector)."""
    if use is not None:
        try:
            code, vec = _link_code_flips(L)
        except Exception:  # noqa: BLE001
            code, vec = nd["dt_code"], nd.get("carry")
        active = True
    else:
        try:
            code = _normalize_dt_code(L.DT_code())
        except Exception:  # noqa: BLE001
            code = nd["dt_code"]
        vec, active = E._own_vector(L, code), False
        if _gate_flags(code):
            try:
                code, vec = _link_code_flips(L)
                active = True
            except Exception:  # noqa: BLE001
                pass
    nd["dt_code"] = code
    nd["n_crossings"] = n
    nd["carry"] = list(vec) if vec is not None else None
    nd["active"] = bool(active and vec is not None)
    if nd["active"] and vec and _gate_flags(code):
        nd["flips"] = list(vec)
        nd["dt_str"] = E.dt_flips_to_string(code, vec)
    else:
        nd.pop("flips", None)
        nd["dt_str"] = _dt_code_to_str(code)


_LAURENT_TERM = __import__("re").compile(
    r"(?P<sign>[+-]?)\s*(?P<coeff>\d+)?\s*(?P<mul>\*)?\s*"
    r"(?P<var>[A-Za-z])?(?:\^\(?(?P<exp>-?\d+(?:/\d+)?)\)?)?")


def _parse_laurent(text):
    """Parse a Laurent-polynomial string into ``{Fraction(exp): int(coeff)}``.

    Returns None if the string cannot be parsed cleanly (so the caller can fall
    back to comparing the raw text).
    """
    from fractions import Fraction
    s = str(text).replace(" ", "")
    if s in ("", "0"):
        return {}
    terms = {}
    consumed = 0
    for m in _LAURENT_TERM.finditer(s):
        if m.start() != consumed:      # a gap means we failed to tokenise
            return None
        piece = m.group(0)
        if piece == "":
            continue
        consumed = m.end()
        sign = -1 if m.group("sign") == "-" else 1
        coeff = m.group("coeff")
        var = m.group("var")
        exp = m.group("exp")
        if coeff is None and var is None:
            return None                # a lone sign etc. -> give up
        c = sign * (int(coeff) if coeff is not None else 1)
        if var is None:
            e = Fraction(0)            # constant term
        elif exp is None:
            e = Fraction(1)            # bare variable
        else:
            e = Fraction(exp)
        terms[e] = terms.get(e, 0) + c
    if consumed != len(s):
        return None
    return {e: c for e, c in terms.items() if c != 0}


def _canonical_jones_key(jones_str):
    """A merge key for a Jones polynomial that is INVARIANT under both

      * the mirror symmetry  q -> 1/q   (negate every exponent), and
      * an overall monomial factor  q^n (shift every exponent).

    The Jones polynomial of a link's mirror image is V(1/q); an overall q^n
    factor is a framing / writhe (Reidemeister-I) normalization artifact rather
    than a topological difference.  Treating both as "the same structure", the
    key normalizes each candidate by shifting its lowest exponent to 0 (dividing
    out q^n) and takes the smaller of the polynomial and its exponent-negated
    (mirror) image.  Coefficients -- including their signs -- are preserved.
    """
    s = str(jones_str).strip()
    if s in ("", "None", "n/a"):
        return ("raw", s)
    d = _parse_laurent(s)
    if not d:                         # None (unparseable) or {} (zero)
        return ("raw", s) if d is None else ("poly", ())

    def _shift_to_zero(pairs):
        mn = min(e for e, _ in pairs)
        return tuple(sorted((e - mn, c) for e, c in pairs))

    orig = _shift_to_zero([(e, c) for e, c in d.items()])
    mirror = _shift_to_zero([(-e, c) for e, c in d.items()])
    return ("poly", min(orig, mirror))


def _short_text(s, n=60):
    s = str(s)
    return s if len(s) <= n else s[:n - 1] + "…"


def _outcome_style(node, original_crossings):
    """Border/fill colours + a short tag describing a node vs. the original."""
    ncr = node["n_crossings"]
    if ncr <= 0:
        return ("#d97706", "#fffbeb", "unknot / unlink")
    if original_crossings is None:
        return ("#6b7280", "#f9fafb", "")
    if ncr < original_crossings:
        return ("#16a34a", "#f0fdf4", "-%d crossings" % (original_crossings - ncr))
    if ncr > original_crossings:
        return ("#ea580c", "#fff7ed", "+%d crossings" % (ncr - original_crossings))
    return ("#6b7280", "#f9fafb", "same crossings")


def _expand_overview_diagram_label_boxes(ax):
    """Give helper-rendered DT/crossing label boxes room in editable SVG."""
    for text_artist in getattr(ax, "texts", []):
        try:
            text_artist.set_fontfamily(OVERVIEW_FONT_FAMILY)
        except Exception:  # noqa: BLE001
            pass
        try:
            bbox_patch = text_artist.get_bbox_patch()
        except Exception:  # noqa: BLE001
            bbox_patch = None
        if bbox_patch is None:
            continue
        try:
            style_name = type(bbox_patch.get_boxstyle()).__name__.lower()
            if "circle" in style_name:
                bbox_patch.set_boxstyle(
                    "circle,pad=%s" % OVERVIEW_DIAGRAM_CROSSING_ID_BOX_PAD)
            else:
                bbox_patch.set_boxstyle(
                    "round,pad=%s" % OVERVIEW_DIAGRAM_DT_LABEL_BOX_PAD)
        except Exception:  # noqa: BLE001
            pass


# Distinct, readable arrow colours (cycled per edge); each arrow's label is
# drawn in the same colour so they are easy to pair up.
ARROW_COLORS = [
    "#e6194B", "#3cb44b", "#4363d8", "#f58231", "#911eb4", "#008080",
    "#f032e6", "#9A6324", "#800000", "#808000", "#000075", "#d2691e",
    "#1f9e89", "#a12f8b", "#2f6f57", "#b5177e", "#6a3d9a", "#b15928",
]


def render_overview_svg(nodes, edges, out_path, negative_even="over",
                        title_dt=None, backtrack_rounds=None,
                        backtrack_steps=None, crossing_order=None,
                        crossing_map=None, crossing_input=None,
                        drawing_options=None, title_has_flips=False):
    """Draw the two-step passage tree as one large, tidy, informative SVG.

    Columns are passage depth (0 = original, 1 = after one passage, 2 = after
    two).  Structures are merged ONLY WITHIN a step, so every column shows its
    own structures.  Each node is a colour-coded card with the helper's default
    2-D view, its DT code, crossing/component counts (blank-line separated), a
    component colour key and an outcome tag; a badge shows how many passages
    arrive at it.  Arrows show operation order; each arrow has its own colour and
    its label (the flipped crossing) is drawn in that same colour, sitting on the
    arrow's arc.  Columns are widely separated and diagrams are enlarged so DT
    and crossing labels do not collide; cards never overlap and diagrams are
    drawn on top of the arrows so nothing is hidden.

    V4.3: a node carrying ``"flips"`` is drawn as that realisation (its
    ``dt_str`` caption is in the flips spelling); ``title_has_flips`` keeps the
    header's start code whole (wrapped, never cut at 66 characters, which
    would drop the vector).
    """
    import textwrap
    import matplotlib as mpl
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

    # Use the same common font while measuring/drawing and while saving.  If
    # the final SVG asks Illustrator for Arial but Matplotlib measured with a
    # different font, editable text can look slightly too large for its boxes.
    mpl.rcParams["font.family"] = OVERVIEW_FONT_FAMILY
    mpl.rcParams["font.sans-serif"] = [OVERVIEW_FONT_FAMILY]
    mpl.rcParams["font.monospace"] = [OVERVIEW_FONT_FAMILY]

    depths = [0, 1, 2]
    by_depth = {d: sorted((n for n in nodes if n["depth"] == d),
                          key=lambda x: (x["n_crossings"], x["id"]))
                for d in depths}
    max_rows = max((len(by_depth[d]) for d in depths), default=1) or 1

    original_crossings = None
    for n in nodes:
        if n["depth"] == 0:
            original_crossings = n["n_crossings"]
            break

    # how many passages arrive at each node (for the merge badge)
    incoming = {}
    for e in edges:
        incoming.setdefault(e["dst"], 0)
        incoming[e["dst"]] += len(e["labels"])

    # Larger canvas + widely separated columns so arrows are legible and the
    # enlarged diagrams have room for their DT/crossing labels.  Figure height
    # scales generously with the tallest column so every (uniformly sized) card
    # is tall enough for a large, uncrowded diagram.
    fig_w = 24.0
    fig_h = max(14.0, 7.4 * max_rows + 3.0)
    fig = plt.figure(figsize=(fig_w, fig_h), facecolor="white")
    overlay = fig.add_axes([0, 0, 1, 1])
    overlay.set_xlim(0, 1)
    overlay.set_ylim(0, 1)
    overlay.axis("off")

    # simplification-parameter string for the title
    if backtrack_rounds and int(backtrack_rounds) > 0:
        sim_info = ("simplify: SnapPy simplify('global') + backtrack  %d rounds "
                    "x %d steps" % (int(backtrack_rounds), int(backtrack_steps or 0)))
    else:
        sim_info = "simplify: SnapPy simplify('global')  (no backtrack)"

    # ---- header band ---------------------------------------------------------
    overlay.add_patch(FancyBboxPatch(
        (0.006, 0.949), 0.988, 0.045, boxstyle="round,pad=0.004",
        linewidth=0, facecolor="#0f172a", zorder=1))
    overlay.text(0.016, 0.978, "Strand-passage overview  -  two steps",
                 ha="left", va="center", color="white", fontsize=14,
                 fontweight="bold", zorder=2)
    overlay.text(0.016, 0.960, sim_info, ha="left", va="center",
                 color="#93c5fd", fontsize=8.5,
                 family=OVERVIEW_FONT_FAMILY, zorder=2)
    if title_dt and title_has_flips:
        title_lines = textwrap.wrap("start  %s" % title_dt, width=120) or [""]
        overlay.text(0.986, 0.978 if len(title_lines) == 1 else 0.990,
                     "\n".join(title_lines), ha="right",
                     va="center" if len(title_lines) == 1 else "top",
                     color="#cbd5e1", fontsize=9 if len(title_lines) == 1 else 7.5,
                     family=OVERVIEW_FONT_FAMILY, zorder=2)
    elif title_dt:
        overlay.text(0.986, 0.978, "start  %s" % _short_text(title_dt, 66),
                     ha="right", va="center", color="#cbd5e1", fontsize=9,
                     family=OVERVIEW_FONT_FAMILY, zorder=2)
    overlay.text(0.986, 0.960, "%d structures, %d passages"
                 % (len(nodes), len(edges)), ha="right", va="center",
                 color="#cbd5e1", fontsize=8.5,
                 family=OVERVIEW_FONT_FAMILY, zorder=2)

    # ---- column banners + geometry (wide separation) -------------------------
    banners = {0: ("Original", "#1d4ed8"),
               1: ("After 1st strand passage", "#0f766e"),
               2: ("After 2nd strand passage", "#7c3aed")}
    col_cx = {0: 0.155, 1: 0.5, 2: 0.845}
    box_w = 0.30
    top, bottom = 0.885, 0.05
    for d in depths:
        label, colour = banners[d]
        overlay.add_patch(FancyBboxPatch(
            (col_cx[d] - box_w / 2, 0.905), box_w, 0.03,
            boxstyle="round,pad=0.004", linewidth=0, facecolor=colour,
            zorder=1))
        overlay.text(col_cx[d], 0.92, "%s   (%d structure%s)"
                     % (label, len(by_depth[d]),
                        "" if len(by_depth[d]) == 1 else "s"),
                     ha="center", va="center", color="white", fontsize=10.5,
                     fontweight="bold", zorder=2)

    # ---- node positions ------------------------------------------------------
    # All cards share ONE fixed size (so every diagram is the same size), and
    # each column's block of cards is centred vertically, so shorter columns sit
    # in the middle and the figure looks symmetric.
    node_pos = {}   # id -> (cx, card_cy, card_w, card_h)
    row_pitch = (top - bottom) / max(max_rows, 1)
    card_h = row_pitch * 0.96
    v_center = 0.5 * (top + bottom)
    for d in depths:
        col = by_depth[d]
        m = len(col)
        if m == 0:
            continue
        block_top = v_center + (m * row_pitch) / 2.0
        for i, n in enumerate(col):
            card_cy = block_top - (i + 0.5) * row_pitch
            node_pos[n["id"]] = (col_cx[d], card_cy, box_w, card_h)

    # ---- arrows first (so diagrams, added later, sit on top) -----------------
    # (Arrows are drawn AFTER the cards, on a top layer, so they are never hidden
    #  under the panels -- see below.)

    # ---- cards + diagrams + captions ----------------------------------------
    for n in nodes:
        if n["id"] not in node_pos:
            continue
        cx, cy, bw, bh = node_pos[n["id"]]
        border, fill, tag = _outcome_style(n, original_crossings)
        if n["depth"] == 0:
            border, fill = "#1d4ed8", "#eff6ff"
        overlay.add_patch(FancyBboxPatch(
            (cx - bw / 2, cy - bh / 2), bw, bh,
            boxstyle="round,pad=0.006", linewidth=1.6,
            edgecolor=border, facecolor=fill, zorder=3))

        # outcome tag (top-left) + merge badge (top-right)
        if tag:
            overlay.text(cx - bw / 2 + 0.008, cy + bh / 2 - bh * 0.035, tag,
                         ha="left", va="top", fontsize=8.0, color=border,
                         fontweight="bold", zorder=6)
        if incoming.get(n["id"], 0) > 1:
            overlay.text(cx + bw / 2 - 0.008, cy + bh / 2 - bh * 0.035,
                         "%d passages merge here" % incoming[n["id"]],
                         ha="right", va="top", fontsize=7.2, color="#7c2d12",
                         style="italic", zorder=6)

        # Big diagram: top ~66% of the card, nearly full width, so the DT and
        # crossing labels are large and uncrowded.
        dw, dh = bw * 0.97, bh * 0.66
        ax = fig.add_axes([cx - dw / 2, cy - bh / 2 + bh * 0.32, dw, dh])
        ax.set_zorder(5)
        try:
            if n.get("flips") is not None:
                # V4.3: every node of a flips run is in spherogram's 'over'
                # convention (run_nongui rewrites an 'under' root)
                diag = E.Diagram.from_dt_code(n["dt_code"], negative_even="over",
                                              flips=n["flips"])
            else:
                diag = E.Diagram.from_dt_code(n["dt_code"],
                                              negative_even=negative_even)
            _apply_crossing_display_options(
                diag, crossing_order, crossing_map, strict=False,
                crossing_input=crossing_input)
            E.render(diag, ax, show_crossing_ids=True, show_dt_labels=True,
                     drawing_options=drawing_options)
            _expand_overview_diagram_label_boxes(ax)
        except Exception as exc:  # noqa: BLE001
            ax.axis("off")
            ax.text(0.5, 0.5, "draw error:\n%s" % exc, ha="center",
                    va="center", fontsize=6, transform=ax.transAxes)

        # Properties (larger, readable font), a blank line between each line.
        dt_wrapped = textwrap.fill(n["dt_str"], width=42)
        jones_wrapped = textwrap.fill("Jones: %s" % n["jones"], width=42)
        cap = "%s\n\ncrossings: %d      components: %d\n\n%s" % (
            dt_wrapped, n["n_crossings"], n["n_components"], jones_wrapped)
        overlay.text(cx, cy - bh / 2 + bh * 0.30, cap, ha="center", va="top",
                     fontsize=8.6, zorder=6, family=OVERVIEW_FONT_FAMILY,
                     color="#111827",
                     linespacing=1.3)

    # ---- arrows LAST, on a dedicated top layer, as STRAIGHT lines ------------
    # A separate full-figure axes added after every card/diagram guarantees the
    # arrows sit on top of everything.  Straight lines stay inside the empty
    # column gaps (a curved arc would bulge sideways into the panels and vanish
    # under them), so nothing is hidden.
    arrow_ax = fig.add_axes([0, 0, 1, 1])
    arrow_ax.set_xlim(0, 1)
    arrow_ax.set_ylim(0, 1)
    arrow_ax.axis("off")
    arrow_ax.set_zorder(20)
    arrow_ax.patch.set_visible(False)

    id2depth = {n["id"]: n["depth"] for n in nodes}
    from collections import defaultdict as _dd0
    gap_groups = _dd0(list)
    for idx, e in enumerate(edges):
        if e["src"] in node_pos and e["dst"] in node_pos:
            gap_groups[id2depth.get(e["src"], 0)].append((idx, e))

    for _gap, items in gap_groups.items():
        # order arrows by target height so neighbouring labels get different t's
        items = sorted(items, key=lambda ie: node_pos[ie[1]["dst"]][1])
        m = len(items)
        for j, (idx, e) in enumerate(items):
            colour = ARROW_COLORS[idx % len(ARROW_COLORS)]
            sx, sy, sbw, sbh = node_pos[e["src"]]
            tx, ty, tbw, tbh = node_pos[e["dst"]]
            start = (sx + sbw / 2, sy)          # source right edge
            end = (tx - tbw / 2, ty)            # target left edge
            arrow_ax.add_patch(FancyArrowPatch(
                start, end, arrowstyle="-|>", mutation_scale=18, lw=2.0,
                color=colour, zorder=21, shrinkA=2, shrinkB=2))
            # label placed ON the straight line, at a t unique to this arrow in
            # the gap, so labels spread out and never stack or overrun panels
            t = 0.30 + 0.40 * (j / (m - 1)) if m > 1 else 0.5
            lx = start[0] + t * (end[0] - start[0])
            ly = start[1] + t * (end[1] - start[1])
            text = "flip " + ", ".join(e["labels"])
            if len(text) > 12:                  # wrap long labels
                text = textwrap.fill(text, width=11)
            arrow_ax.text(lx, ly, text, ha="center", va="center", fontsize=7.4,
                          zorder=22, color=colour, fontweight="bold",
                          bbox=dict(boxstyle="round,pad=%s"
                                    % OVERVIEW_TEXT_BOX_PAD, fc="white",
                                    ec=colour, lw=1.1, alpha=0.97))

    # ---- footer legend -------------------------------------------------------
    legend = ("Card colour vs. original crossing count:  "
              "green = fewer  |  gray = same  |  orange = more  |  "
              "gold = unknot/unlink.    "
              "Each step is deduplicated on its own (structures are NOT merged "
              "across steps); within a step, structures merge when they share "
              "crossing count, component count and Jones up to mirror (q->1/q) "
              "and framing (overall q^n).    "
              "Each arrow has its own colour; its label (flipped DT crossing) is "
              "drawn in the same colour.")
    overlay.text(0.5, 0.026, legend, ha="center", va="center", fontsize=8.0,
                 color="#374151", zorder=4, wrap=True,
                 bbox=dict(boxstyle="round,pad=0.65",
                           fc="#f3f4f6", ec="#d1d5db"))

    # Illustrator-friendly SVG: keep labels editable as real SVG <text> while
    # using Arial throughout the overview to avoid most font-substitution
    # warnings.  Artist clipping stays disabled so Illustrator does not warn
    # about clipping loss on SVG roundtrip.
    for artist in fig.findobj():
        try:
            artist.set_clip_on(False)
        except Exception:  # noqa: BLE001
            pass
        try:
            if isinstance(artist, mpl.text.Text):
                artist.set_fontfamily(OVERVIEW_FONT_FAMILY)
        except Exception:  # noqa: BLE001
            pass
    with mpl.rc_context({
            "svg.fonttype": "none",
            "font.family": OVERVIEW_FONT_FAMILY,
            "font.sans-serif": [OVERVIEW_FONT_FAMILY],
            "font.monospace": [OVERVIEW_FONT_FAMILY],
    }):
        fig.savefig(out_path, format="svg", facecolor="white")
    plt.close(fig)
    # V4.2: Matplotlib writes the font as a CSS shorthand that Illustrator does
    # not reliably honour; restate it as longhand plus a fallback stack.
    _rewrite_svg_font_declarations(out_path)


def run_nongui(dt_string, out_path, negative_even="over",
               backtrack_rounds=0, backtrack_steps=20,
               crossing_order=None, crossing_map=None, crossing_input=None,
               drawing_options=None, drawing_session_path=None):
    """Two-pass strand-passage spreadsheet plus a merged overview SVG.

    ``backtrack_rounds`` > 0 (V3.4) applies backtrack-assisted simplification to
    every SnapPy simplification here (original link and every passage result).

    V3.5 enumerates the second pass once per merged/reconciled first-step
    structure, not once per raw first-step crossing.

    V3.6 continues only first-step structures with ``new_components > 2`` and a
    usable chosen DT code.

    V4.3: ``dt_string`` may carry a flip vector ('DT: [...], [0,1,...]'); then
    the given realisation is followed through both passes (see
    ``strand_passage_nongui``), every node keeps its vector, and the workbook
    and overview show them.  A code read with ``negative_even='under'`` --
    with or without a vector -- is first rewritten in spherogram's 'over'
    convention (entries negated, the same diagram; a vector describes the
    shadow and is unchanged), which every SnapPy step reads and writes, so
    every code in the workbook and on the overview cards is 'over' and each
    card is drawn as the code SnapPy read (an [info] line and the run_info
    row ``negative_even_note`` say so for a bare code).  Otherwise, without a
    vector: exactly as before.
    """
    try:
        import snappy  # type: ignore
    except Exception as exc:  # noqa: BLE001
        print("[error] SnapPy is required for --nongui. Run under "
              "'sage -python'. (%s)" % exc, file=sys.stderr)
        return 2
    try:
        import pandas as pd  # type: ignore
    except Exception as exc:  # noqa: BLE001
        print("[error] pandas is required for --nongui (%s)" % exc,
              file=sys.stderr)
        return 2

    drawing_options = normalize_drawing_options(drawing_options)
    try:
        crossing_order, crossing_map, detected_crossing_kind = _split_crossing_display_input(
            crossing_input=crossing_input,
            crossing_order=crossing_order,
            crossing_map=crossing_map,
        )
    except Exception as exc:  # noqa: BLE001
        print("[error] invalid crossing-label input: %s" % exc, file=sys.stderr)
        return 2

    try:
        dt_code, flips = E.parse_dt_flips(dt_string)
    except ValueError as exc:
        if not _text_carries_vector(dt_string):
            # round 4d (2026-10-20): a bare code the parser cannot read
            # ('DT[daddcab]', '8,6,2,4') is ONE [error] line, exit 2 -- it
            # was a traceback (with a per-process object address in it)
            print("[error] %s" % _unparseable_text(dt_string, exc), file=sys.stderr)
            return 2
        print("[error] could not read the DT code: %s" % exc, file=sys.stderr)
        return 2
    flips_note = None
    under_note = None
    given = flips is not None       # the user named the realisation
    if flips is not None:
        if negative_even == "under":
            # every SnapPy step reads spherogram's 'over' convention; the same
            # diagram there is the negated code, with the same vector
            dt_code = [tuple(-x for x in comp) for comp in dt_code]
            flips_note = ("--negative-even under: the code was rewritten in "
                          "spherogram's 'over' convention (entries negated; "
                          "the vector describes the shadow and is unchanged)")
            print("[info] %s: %s" % (flips_note, E.dt_flips_to_string(dt_code, flips)))
        try:
            E.link_from_dt(snappy, dt_code, flips)
        except ValueError as exc:
            print("[error] %s" % exc, file=sys.stderr)
            return 2
        print("[info] flip vector %s: following this realisation through "
              "every passage" % _flip_vector_text(flips))
    elif negative_even == "under":
        # V4.3: SnapPy reads spherogram's 'over' convention and writes its
        # codes in it, so the bare code is rewritten too (the same diagram);
        # every code in the workbook and on the overview cards is then 'over'
        dt_code = [tuple(-x for x in comp) for comp in dt_code]
        under_note = ("--negative-even under: the code was rewritten in "
                      "spherogram's 'over' convention (entries negated, the "
                      "same diagram), which every SnapPy step reads; every DT "
                      "code below is in that convention")
        print("[info] %s: %s" % (under_note, _dt_code_to_str(dt_code)))
    root_carry, root_active = None, False
    if flips is None:
        # V4.3 (2026-10-19, round 3): every bare root carries its realisation
        # -- spherogram's reading (dt_converter's mended one for a flagged
        # code it reads illegally) -- through both passes; it is in force,
        # and written out, where the composite gate flags a code (one [info]
        # line when that is the root itself)
        root_carry, flagged, line = _bare_root_realisation(dt_code, "over")
        followed = False
        if (not flagged or not root_carry) and _bare_root_unbuildable(snappy, dt_code):
            # round 4 (2026-10-19): spherogram cannot build the bare root (its
            # decoder fails, or would loop for ever on a kink: DecoderLoop),
            # and every step below read it with snappy.Link -- a traceback, or
            # no return at all.  Follow a reading of it instead (a code the
            # gate does not flag is the same link up to mirror image in every
            # realisation), or stop in one line where there is none
            if not any(dt_code):
                print("[error] %s has no crossings (the unknot or an unlink), "
                      "so --nongui has no strand passage to make"
                      % _dt_code_to_str(dt_code), file=sys.stderr)
                return 2
            mended, _flagged, mline = _bare_root_realisation(
                dt_code, "over", follow_unreadable=True)
            if mended and mline is None:
                # spherogram's own reading, which its string parser cannot build
                mline = ("[info] spherogram cannot build the bare code %s (%s), "
                         "so the run follows spherogram's own reading of it, %s, "
                         "through every step (the composite gate does not flag "
                         "it, so every realisation of it is the same link up to "
                         "mirror image; the root is written with this vector, "
                         "every other code the gate does not flag bare)"
                         % (_dt_code_to_str(dt_code), _bare_root_unbuildable(snappy, dt_code),
                            E.dt_flips_to_string(dt_code, mended)))
            if not mended or (mline or "").startswith("[error]"):
                print(mline or ("[error] spherogram cannot build the bare code %s "
                                "(%s), and no reading of it was found; give the "
                                "code with a flip vector"
                                % (_dt_code_to_str(dt_code),
                                   _bare_root_unbuildable(snappy, dt_code))),
                      file=sys.stderr)
                return 2
            root_carry, line, followed = mended, mline, not flagged
        if flagged and root_carry:
            try:
                E.link_from_dt(snappy, dt_code, root_carry)
            except ValueError as exc:
                line = ("[warn] the composite gate flags the bare code, but "
                        "spherogram cannot build its own reading of it (%s), "
                        "so it is read bare as before" % exc)
                root_carry = None
        if line is not None:
            print(line, file=sys.stderr if line.startswith("[warn]") else sys.stdout)
        root_active = bool((flagged or followed) and root_carry)
        if followed:
            # round 4: an unflagged root spherogram cannot build, followed
            flips_note = (
                "the flip vector the run follows for the bare code, which "
                "spherogram cannot build bare (the run's [info] line says why "
                "and whose reading this is); the composite gate does not flag "
                "it, so every realisation is the same link up to mirror image; "
                "bit k <-> the k-th DT entry (odd label 2k+1); walking along the "
                "visit with the smaller label, 0 = the other strand crosses from "
                "right to left (spherogram/SnapPy convention); every link here "
                "follows (DT, flips); the root is written with this vector, every "
                "other code with its vector only where the gate flags it")
        elif root_active:
            flips_note = (
                "%s of the bare code (the composite gate flagged it: a bare "
                "code does not fix how its parts are reflected against each "
                "other, so re-reading it could name another link); bit k <-> "
                "the k-th DT entry (odd label 2k+1); walking along the visit "
                "with the smaller label, 0 = the other strand crosses from "
                "right to left (spherogram/SnapPy convention); every link "
                "here follows (DT, flips), and the codes the gate flags are "
                "written with their vectors"
                % ("dt_converter's mended reading" if "mended reading" in line
                   else "dt_converter's reading" if "dt_converter's reading" in line
                   else "a legal reading" if "a legal reading" in line
                   else "spherogram's reading"))
    # the realisation the workbook names for the root (the user's vector, or
    # the carried one of a flagged bare root)
    shown_root_flips = flips if given else (root_carry if root_active else None)
    # the convention every code from here on is read in
    read_as = "over" if negative_even == "under" else negative_even
    if (crossing_order and crossing_order.strip()) or (crossing_map and crossing_map.strip()):
        try:
            if given:
                root_for_ids = E.Diagram.from_dt_code(dt_code, negative_even="over",
                                                      flips=flips)
            else:
                root_for_ids = E.Diagram.from_dt_code(dt_code, negative_even=read_as)
            _apply_crossing_display_options(
                root_for_ids, crossing_order, crossing_map, strict=True)
        except Exception as exc:  # noqa: BLE001
            print("[error] invalid crossing-label input: %s" % exc,
                  file=sys.stderr)
            return 2
        print("[info] custom displayed crossing IDs detected as %s"
              % detected_crossing_kind)

    # Merged-structure node/edge accumulator for the overview SVG.
    nodes: Dict[Any, Dict[str, Any]] = {}
    order: List[Any] = []
    edges: Dict[Tuple[int, int], set] = {}

    def get_node(code, depth, simpl_crossings, ncomp, jones, node_flips=None,
                 node_carry=None, node_active=False):
        """Add (or fetch) a node for one step.

        Structures are merged ONLY WITHIN THE SAME STEP: the merge key includes
        the depth, so a structure that recurs at a later step (even if it equals
        the original) gets its own card in its own column.  Within a step, two
        structures merge when they share the SnapPy-simplified crossing count,
        the component count, and the Jones polynomial up to mirror (q -> 1/q) and
        framing (overall q^n).  The crossing/component/Jones values are the ones
        already computed for the spreadsheet, so the two outputs always agree.

        V4.3: ``node_flips`` (flips runs) is stored as the node's ``"flips"``
        and its ``dt_str`` is in the flips spelling.  Round 3: a node of a
        bare run also keeps ``"carry"`` / ``"active"`` (the realisation it
        carries, unseen unless its code is flagged; ``node_flips`` is then
        set too).
        """
        code = [tuple(int(x) for x in comp) for comp in code]
        simpl = simpl_crossings if simpl_crossings is not None else _crossing_count(code)
        if node_flips is not None:
            comp_n = (ncomp if ncomp is not None
                      else _num_components(snappy, code, node_flips))
        elif node_active and node_carry:
            comp_n = (ncomp if ncomp is not None
                      else _num_components(snappy, code, node_carry))
        else:
            comp_n = ncomp if ncomp is not None else _num_components(snappy, code)
        fp = (int(depth), int(simpl), int(comp_n), _canonical_jones_key(jones))
        if fp not in nodes:
            nid = len(order)
            nodes[fp] = {
                "id": nid,
                "dt_code": code,
                "dt_str": _dt_code_to_str(code),
                "n_crossings": _crossing_count(code),
                "n_components": comp_n,
                "jones": str(jones),
                "depth": depth,
                "fp": fp,
            }
            if node_flips is not None:
                nodes[fp]["flips"] = list(node_flips)
                nodes[fp]["dt_str"] = E.dt_flips_to_string(code, node_flips)
            if not given:
                nodes[fp]["carry"] = (list(node_carry)
                                      if node_carry is not None else None)
                nodes[fp]["active"] = bool(node_active and node_carry is not None)
            order.append(fp)
        return nodes[fp]["id"]

    if backtrack_rounds:
        print("[info] backtrack-assisted simplification ON (%d rounds, %d steps)"
              % (backtrack_rounds, backtrack_steps))

    # Original-link invariants, computed the same way the spreadsheet does.
    if given:
        try:
            L0 = E.link_from_dt(snappy, dt_code, flips)
            L0 = E.backtrack_simplify(snappy, L0, mode='global',
                                      rounds=backtrack_rounds,
                                      steps=backtrack_steps, archive_flips=True)
            root_crossings = len(L0.crossings)
            root_components = len(L0.link_components)
            root_jones = _jones_for_dt(snappy, dt_code, flips)
        except Exception:  # noqa: BLE001
            root_crossings = _crossing_count(dt_code)
            root_components = _num_components(snappy, dt_code, flips)
            root_jones = "None"
        root_id = get_node(dt_code, 0, root_crossings, root_components,
                           root_jones, node_flips=flips)
    elif root_active:
        # round 3: a flagged bare root, followed as it carries
        try:
            L0 = E.link_from_dt(snappy, dt_code, root_carry)
            L0 = E.backtrack_simplify(snappy, L0, mode='global',
                                      rounds=backtrack_rounds,
                                      steps=backtrack_steps, archive_flips=True)
            root_crossings = len(L0.crossings)
            root_components = len(L0.link_components)
            root_jones = _jones_for_dt(snappy, dt_code, root_carry)
        except Exception:  # noqa: BLE001
            root_crossings = _crossing_count(dt_code)
            root_components = _num_components(snappy, dt_code, root_carry)
            root_jones = "None"
        root_id = get_node(dt_code, 0, root_crossings, root_components,
                           root_jones, node_flips=root_carry,
                           node_carry=root_carry, node_active=True)
    else:
        try:
            L0 = _bare_snappy_link(snappy, _dt_code_to_str(dt_code))
            L0 = E.backtrack_simplify(snappy, L0, mode='global',
                                      rounds=backtrack_rounds, steps=backtrack_steps,
                                      archive_flips="auto")
            root_crossings = len(L0.crossings)
            root_components = len(L0.link_components)
            root_jones = _jones_for_dt(snappy, dt_code)
        except Exception:  # noqa: BLE001
            root_crossings = _crossing_count(dt_code)
            root_components = _num_components(snappy, dt_code)
            root_jones = "None"
        root_id = get_node(dt_code, 0, root_crossings, root_components, root_jones,
                           node_carry=root_carry, node_active=False)

    print("[info] first pass ...")
    if given:
        results, chosen_codes, chosen_flips = strand_passage_nongui(
            snappy, dt_code, backtrack_rounds=backtrack_rounds,
            backtrack_steps=backtrack_steps, flips=flips)
        chosen_meta = [{"flips": f, "carry": None, "active": False}
                       for f in chosen_flips]
    else:
        results, chosen_codes, chosen_meta = strand_passage_nongui(
            snappy, dt_code, backtrack_rounds=backtrack_rounds,
            backtrack_steps=backtrack_steps, carry=root_carry,
            carry_active=root_active, return_meta=True)

    print("[info] first-step merge ...")
    print("[info] second-pass continuation criterion: %s"
          % NONGUI_SECOND_PASS_CRITERION)
    second_pass_results_dict = {}
    continuable_by_label: Dict[str, Dict[str, Any]] = {}
    continuable_first_passages = 0
    # V4.1: remember which merged structure each spreadsheet row landed on, so
    # the row can be shaded with the colour of the card that structure is drawn
    # as.  The ids recorded here are pre-merge; they are chased through every
    # later remap by ``_follow`` below.
    first_row_nid: List[Optional[int]] = []
    for res, chosen_code, cm in zip(results, chosen_codes, chosen_meta):
        chosen_fl = cm["flips"]
        if chosen_fl is not None:
            nid = get_node(chosen_code, 1, res['snappy_crossings'],
                           res['new_components'], res['Jones_polynomial'],
                           node_flips=chosen_fl, node_carry=cm["carry"],
                           node_active=cm["active"])
        else:
            nid = get_node(chosen_code, 1, res['snappy_crossings'],
                           res['new_components'], res['Jones_polynomial'],
                           node_carry=cm["carry"], node_active=cm["active"])
        first_row_nid.append(nid)
        label = str(res['flipped_crossing'])
        edges.setdefault((root_id, nid), set()).add(label)
        new_components = res.get('new_components')
        if new_components is not None and int(new_components) > 2 and chosen_code:
            continuable_first_passages += 1
            continuable_by_label[label] = res

    # V3.5: before enumerating the second pass, collapse the first-step graph to
    # the same merged representatives shown in the overview.  That prevents
    # equivalent first-pass diagrams from spawning duplicate second-pass sheets.
    first_node_list = [nodes[fp] for fp in order]
    first_edge_list = _edge_list_from_map(edges)
    remap_first: Dict[int, int] = {}
    if backtrack_rounds:
        before_n = len(first_node_list)
        first_node_list, first_edge_list, remap_first = reconcile_steps(
            snappy, first_node_list, first_edge_list,
            rounds=max(300, int(backtrack_rounds)), steps=backtrack_steps)
        if len(first_node_list) < before_n:
            print("[info] first-step reconciliation merged %d redundant "
                  "structure(s) before continuation"
                  % (before_n - len(first_node_list)))
    first_node_list, first_edge_list, renum_first = _renumber_passage_graph(
        first_node_list, first_edge_list)

    def _follow(node_id, *id_maps):
        """Chase a node id through the remaps applied after it was created."""
        current = node_id
        for id_map in id_maps:
            if current is None:
                return None
            current = id_map.get(current, current)
        return current

    first_row_nid = [_follow(n, remap_first, renum_first) for n in first_row_nid]

    # Rebuild the node/edge accumulators from the merged first-step graph, then
    # append second-step nodes below using the same get_node() merge key.
    nodes.clear()
    order[:] = []
    edges.clear()
    for nd in first_node_list:
        nodes[nd["fp"]] = nd
        order.append(nd["fp"])
    for edge in first_edge_list:
        edges[(edge["src"], edge["dst"])] = set(edge["labels"])

    root_id = next((nd["id"] for nd in first_node_list if nd["depth"] == 0),
                   root_id)
    id_to_node = {nd["id"]: nd for nd in first_node_list}
    first_step_labels: Dict[int, set] = {}
    for edge in first_edge_list:
        if edge["src"] != root_id:
            continue
        dst = id_to_node.get(edge["dst"])
        if dst is None or dst["depth"] != 1:
            continue
        labels = [str(x) for x in edge["labels"] if str(x) in continuable_by_label]
        if labels:
            first_step_labels.setdefault(edge["dst"], set()).update(labels)

    if continuable_first_passages:
        print("[info] second pass uses %d merged first-step structure(s) "
              "from %d continuable first-step passage(s)"
              % (len(first_step_labels), continuable_first_passages))
    else:
        print("[info] second pass uses 0 merged first-step structure(s) "
              "from 0 continuable first-step passage(s)")

    used_sheet_names = set()
    second_row_nid: Dict[str, List[Optional[int]]] = {}
    sheet_parent_nid: Dict[str, int] = {}
    print("[info] second pass ...")
    for nid in sorted(first_step_labels,
                      key=lambda x: (id_to_node[x]["n_crossings"], x)):
        node = id_to_node[nid]
        chosen_code = node["dt_code"]
        labels = sorted(first_step_labels[nid], key=_label_sort_key)
        provenance = ", ".join(labels)
        representative = labels[0]
        sheet_name = _safe_sheet_name(
            "merged_%s_%s" % (nid, "_".join(labels[:3])),
            used_sheet_names)
        try:
            node_flips = node.get("flips")
            if given and node_flips is not None:
                second_results, second_chosen, second_flips = strand_passage_nongui(
                    snappy, list(chosen_code),
                    backtrack_rounds=backtrack_rounds,
                    backtrack_steps=backtrack_steps, flips=node_flips)
                second_meta = [{"flips": f, "carry": None, "active": False}
                               for f in second_flips]
            else:
                # round 3: a bare run's node continues with the realisation
                # it carries (in force when its lineage met a flagged code)
                second_results, second_chosen, second_meta = strand_passage_nongui(
                    snappy, list(chosen_code),
                    backtrack_rounds=backtrack_rounds,
                    backtrack_steps=backtrack_steps, carry=node.get("carry"),
                    carry_active=bool(node.get("active")), return_meta=True)
            for item in second_results:
                item['first_step_node_id'] = nid
                item['first_step_passages'] = provenance
                item['first_step_representative'] = representative
            second_pass_results_dict[sheet_name] = second_results
            sheet_parent_nid[sheet_name] = nid
            sheet_row_nid: List[Optional[int]] = []
            for r2, c2, m2 in zip(second_results, second_chosen, second_meta):
                f2 = m2["flips"]
                if f2 is not None:
                    nid2 = get_node(c2, 2, r2['snappy_crossings'],
                                    r2['new_components'], r2['Jones_polynomial'],
                                    node_flips=f2, node_carry=m2["carry"],
                                    node_active=m2["active"])
                else:
                    nid2 = get_node(c2, 2, r2['snappy_crossings'],
                                    r2['new_components'], r2['Jones_polynomial'],
                                    node_carry=m2["carry"],
                                    node_active=m2["active"])
                sheet_row_nid.append(nid2)
                edges.setdefault((nid, nid2), set()).add(
                    str(r2['flipped_crossing']))
            second_row_nid[sheet_name] = sheet_row_nid
        except Exception as exc:  # noqa: BLE001
            print("[warn] second pass failed for merged first-step node %s "
                  "(passage%s %s): %s"
                  % (nid, "" if len(labels) == 1 else "s",
                     provenance, exc), file=sys.stderr)

    if not out_path:
        out_path = "strand_passage_results.xlsx"
    if not str(out_path).endswith(".xlsx"):
        out_path = str(out_path) + ".xlsx"

    overview_path = out_path[:-5] + "_overview.svg"
    second_pass_rows = sum(len(sheet_data)
                           for sheet_data in second_pass_results_dict.values())

    # ---- final merged graph --------------------------------------------------
    # V4.1: the last reconciliation runs HERE, before the workbook is written.
    # It can both further simplify a structure and merge two of them, so doing it
    # after the write (as V4.0 did) left the spreadsheet describing structures
    # that no card in the figure actually showed.  Both outputs now come from
    # this one final graph.
    node_list = [nodes[fp] for fp in order]
    edge_list = [{"src": s, "dst": d, "labels": sorted(l, key=_label_sort_key)}
                 for (s, d), l in edges.items()]

    # Targeted per-step reconciliation: pull any incompletely-simplified
    # same-Jones structures down to their step-mate's crossing count and merge.
    remap_final: Dict[int, int] = {}
    if backtrack_rounds:
        before_n = len(node_list)
        node_list, edge_list, remap_final = reconcile_steps(
            snappy, node_list, edge_list,
            rounds=max(300, int(backtrack_rounds)), steps=backtrack_steps)
        if len(node_list) < before_n:
            print("[info] reconciliation merged %d redundant structure(s)"
                  % (before_n - len(node_list)))

    # ---- tie every workbook row to the card it is drawn as -------------------
    node_by_id = {nd["id"]: nd for nd in node_list}
    card_sheets: Dict[int, set] = {}

    def _assign_cards(rows, row_nids, sheet_name):
        """Stamp ``svg_card_id`` on each row; returns the per-row card ids."""
        card_ids = []
        for i, row in enumerate(rows):
            nid = row_nids[i] if i < len(row_nids) else None
            card_id = _follow(nid, remap_final)
            if card_id not in node_by_id:
                card_id = None      # errored passage: nothing was drawn for it
            row['svg_card_id'] = card_id
            if card_id is not None:
                card_sheets.setdefault(card_id, set()).add(sheet_name)
            card_ids.append(card_id)
        return card_ids

    first_card_ids = _assign_cards(results, first_row_nid, 'first_pass')
    second_card_ids = {}
    for sheet_name, sheet_data in second_pass_results_dict.items():
        # The parent card id can itself have been merged away by the final
        # reconciliation, so re-point it at the card that survived.
        parent_card = _follow(sheet_parent_nid.get(sheet_name), remap_final)
        for row in sheet_data:
            row['first_step_node_id'] = parent_card
        second_card_ids[sheet_name] = _assign_cards(
            sheet_data, second_row_nid.get(sheet_name, []), sheet_name)

    original_crossings = next((nd["n_crossings"] for nd in node_list
                               if nd["depth"] == 0), None)
    card_index = _card_index_rows(node_list, edge_list, card_sheets,
                                  original_crossings)

    run_info = _run_info_rows(
        snappy, pd, dt_string, dt_code, out_path, overview_path,
        negative_even, backtrack_rounds, backtrack_steps,
        crossing_order, crossing_map, crossing_input, drawing_options,
        drawing_session_path, len(results), continuable_first_passages,
        len(first_step_labels), len(second_pass_results_dict), second_pass_rows,
        **({"flips": shown_root_flips, "flips_note": flips_note}
           if shown_root_flips is not None else {}),
        **({"negative_even_note": under_note} if under_note is not None else {}))
    run_info.extend([
        ("svg_cards_total", len(node_list)),
        ("row_shading", "each data row is shaded with the colour of the "
                        "overview-SVG card it is drawn as; see the card_index "
                        "sheet and the svg_card_id column"),
    ])

    with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
        pd.DataFrame(run_info, columns=["field", "value"]).to_excel(
            writer, index=False, sheet_name='run_info')

        shaded = 0
        if card_index:
            card_df = pd.DataFrame(card_index)
            card_df.to_excel(writer, index=False, sheet_name='card_index')
            _shade_rows_by_card(writer.sheets['card_index'],
                                [r["svg_card_id"] for r in card_index],
                                len(card_df.columns), node_by_id)

        first_df = pd.DataFrame(results)
        first_df.to_excel(writer, index=False, sheet_name='first_pass')
        shaded += _shade_rows_by_card(writer.sheets['first_pass'],
                                      first_card_ids, len(first_df.columns),
                                      node_by_id)
        for sheet_name, sheet_data in second_pass_results_dict.items():
            sheet_df = pd.DataFrame(sheet_data)
            sheet_df.to_excel(writer, index=False, sheet_name=sheet_name)
            shaded += _shade_rows_by_card(writer.sheets[sheet_name],
                                          second_card_ids.get(sheet_name, []),
                                          len(sheet_df.columns), node_by_id)
    print("[ok] wrote %s  (%d card(s); %d row(s) shaded by card)"
          % (out_path, len(node_list), shaded))
    if not given and not root_active:
        flagged_cards = sum(1 for nd in node_list if nd.get("flips") is not None)
        if flagged_cards:
            # round 3: a prime-looking root reached composite codes
            print("[info] the composite gate flagged %d of the %d card(s); "
                  "their codes are written with the flip vectors of the "
                  "realisation followed from the root (dt_flip_vector "
                  "columns), every other code bare" % (flagged_cards,
                                                       len(node_list)))

    try:
        if shown_root_flips is not None:
            # the user's own code text, with its vector in the toolkit's
            # spelling, never cut (the vector would be lost)
            render_overview_svg(node_list, edge_list, overview_path,
                                negative_even="over",
                                title_dt=format_dt_flips(
                                    split_dt_flips(dt_string)[0].strip(),
                                    shown_root_flips),
                                backtrack_rounds=backtrack_rounds,
                                backtrack_steps=backtrack_steps,
                                crossing_order=crossing_order,
                                crossing_map=crossing_map,
                                drawing_options=drawing_options,
                                title_has_flips=True)
        else:
            render_overview_svg(node_list, edge_list, overview_path,
                                negative_even=read_as, title_dt=dt_string,
                                backtrack_rounds=backtrack_rounds,
                                backtrack_steps=backtrack_steps,
                                crossing_order=crossing_order,
                                crossing_map=crossing_map,
                                drawing_options=drawing_options)
        print("[ok] wrote %s  (%d unique structures, %d passage arrows)"
              % (overview_path, len(node_list), len(edge_list)))
    except Exception as exc:  # noqa: BLE001
        print("[warn] could not write overview SVG: %s" % exc, file=sys.stderr)

    return 0


# --------------------------------------------------------------------------- #
#  Matplotlib-in-Tk panel with a TkAgg -> Agg fallback, plus hover cursor
# --------------------------------------------------------------------------- #
def _make_figure_panel(fig, ax, master, tk, ttk, click_callback,
                       hover_test=None, backend="auto"):
    requested = (backend or "auto").lower()
    if requested not in ("auto", "tkagg", "agg"):
        requested = "auto"

    tkagg_error = None
    if requested in ("auto", "tkagg"):
        try:
            from matplotlib.backends.backend_tkagg import (  # type: ignore
                FigureCanvasTkAgg, NavigationToolbar2Tk)
            canvas = FigureCanvasTkAgg(fig, master=master)
            widget = canvas.get_tk_widget()
            widget.pack(fill=tk.BOTH, expand=True)
            try:
                NavigationToolbar2Tk(canvas, master).update()
            except Exception:  # noqa: BLE001
                pass

            def draw_idle():
                canvas.draw_idle()

            def _wrapped(event):
                if event.inaxes is not ax or event.xdata is None:
                    return
                click_callback(float(event.xdata), float(event.ydata))

            def _motion(event):
                if hover_test is None:
                    return
                try:
                    if event.inaxes is not ax or event.xdata is None:
                        widget.configure(cursor="")
                        return
                    near = hover_test(float(event.xdata), float(event.ydata))
                    widget.configure(cursor="hand2" if near else "")
                except Exception:  # noqa: BLE001
                    pass

            fig.canvas.mpl_connect("button_press_event", _wrapped)
            fig.canvas.mpl_connect("motion_notify_event", _motion)
            return {"name": "TkAgg", "draw_idle": draw_idle}
        except Exception as exc:  # noqa: BLE001
            tkagg_error = exc
            if requested == "tkagg":
                raise

    import base64
    import io
    from matplotlib.backends.backend_agg import FigureCanvasAgg

    agg = FigureCanvasAgg(fig)
    frame = ttk.Frame(master)
    frame.pack(fill=tk.BOTH, expand=True)
    msg = "Agg image display"
    if tkagg_error is not None:
        msg += " (TkAgg unavailable: %s)" % tkagg_error
    ttk.Label(frame, text=msg, wraplength=640).pack(anchor="w", pady=(0, 3))
    tk_canvas = tk.Canvas(frame, highlightthickness=0, background="white")
    tk_canvas.pack(fill=tk.BOTH, expand=True)

    def draw_idle():
        agg.draw()
        width, height = agg.get_width_height()
        buf = io.BytesIO()
        agg.print_png(buf)
        data = base64.b64encode(buf.getvalue()).decode("ascii")
        try:
            photo = tk.PhotoImage(data=data)
        except Exception:  # noqa: BLE001
            photo = tk.PhotoImage(data=data, format="png")
        tk_canvas.delete("all")
        tk_canvas.config(scrollregion=(0, 0, width, height),
                         width=width, height=height)
        tk_canvas.create_image(0, 0, anchor="nw", image=photo)
        tk_canvas._sp_photo = photo  # keep a reference alive

    def _agg_click(event):
        width, height = agg.get_width_height()
        if event.x < 0 or event.y < 0 or event.x > width or event.y > height:
            return
        xdata, ydata = ax.transData.inverted().transform((event.x, height - event.y))
        click_callback(float(xdata), float(ydata))

    def _agg_motion(event):
        if hover_test is None:
            return
        width, height = agg.get_width_height()
        if event.x < 0 or event.y < 0 or event.x > width or event.y > height:
            tk_canvas.configure(cursor="")
            return
        try:
            xdata, ydata = ax.transData.inverted().transform(
                (event.x, height - event.y))
            near = hover_test(float(xdata), float(ydata))
            tk_canvas.configure(cursor="hand2" if near else "")
        except Exception:  # noqa: BLE001
            pass

    tk_canvas.bind("<Button-1>", _agg_click)
    tk_canvas.bind("<Motion>", _agg_motion)
    return {"name": "Agg-in-Tk", "draw_idle": draw_idle}


# --------------------------------------------------------------------------- #
#  Windowed application: every passage is a new window
# --------------------------------------------------------------------------- #
class StrandPassageApp(object):
    def __init__(self, root, tk, ttk, dt_string, negative_even="over",
                 use_snappy_global=True, gui_backend="auto",
                 backtrack_enabled=True,
                 backtrack_rounds=DEFAULT_BACKTRACK_ROUNDS,
                 backtrack_steps=DEFAULT_BACKTRACK_STEPS,
                 crossing_order=None, crossing_map=None, crossing_input=None,
                 drawing_options=None, drawing_session_path=None):
        self.root = root
        self.tk = tk
        self.ttk = ttk
        self.negative_even = negative_even
        self.use_snappy_global = bool(use_snappy_global)
        self.gui_backend = gui_backend
        self.init_backtrack_enabled = bool(backtrack_enabled)
        self.init_backtrack_rounds = int(backtrack_rounds)
        self.init_backtrack_steps = int(backtrack_steps)
        order, cmap, _kind = _split_crossing_display_input(
            crossing_input=crossing_input,
            crossing_order=crossing_order,
            crossing_map=crossing_map,
        )
        self.init_crossing_input = cmap or order
        self.drawing_options = normalize_drawing_options(drawing_options)
        self.drawing_session_path = drawing_session_path or ""
        self.windows: List[Any] = []
        self.step_count = 0
        self.tile = TILE

        self._build_root_controls(dt_string)
        self.load()

    def _current_backtrack(self):
        """Return (rounds, steps) from the controls; rounds=0 when disabled."""
        if not self.backtrack_var.get():
            return 0, self._safe_int(self.backtrack_steps_var,
                                     DEFAULT_BACKTRACK_STEPS)
        return (self._safe_int(self.backtrack_rounds_var,
                               DEFAULT_BACKTRACK_ROUNDS),
                self._safe_int(self.backtrack_steps_var,
                               DEFAULT_BACKTRACK_STEPS))

    @staticmethod
    def _safe_int(var, fallback):
        try:
            return max(0, int(float(var.get())))
        except Exception:  # noqa: BLE001
            return fallback

    def _current_crossing_options(self):
        return _split_crossing_display_input(
            crossing_input=self.crossing_input_var.get())

    def _current_drawing_options(self):
        return normalize_drawing_options(self.drawing_options)

    def _show_help(self, title, message):
        self._toast(self.root, message)

    def _show_snappy_help(self):
        self._show_help(
            "SnapPy global",
            "SnapPy global controls whether the GUI asks SnapPy/Sage to "
            "simplify a diagram and compute invariants such as the Jones "
            "polynomial. If it is off, strand passages use only the continuity "
            "engine and the Simplify button is skipped.")

    def _show_backtrack_help(self):
        self._show_help(
            "Backtrack simplify",
            "Backtrack simplify repeatedly complicates a SnapPy link and then "
            "runs simplify('global'), keeping the fewest-crossing result found. "
            "Rounds is the number of attempts; steps is the complication length "
            "inside each attempt. Larger values can find lower plateaus but take "
            "longer.")

    # ---- root control strip ----
    def _build_root_controls(self, dt_string):
        tk, ttk = self.tk, self.ttk
        self.root.title("DT Link Toolkit strand passage V%s  -  click a crossing to "
                        "branch a new window" % VERSION)
        controls = ttk.Frame(self.root, padding=6)
        controls.pack(side=tk.TOP, fill=tk.X)
        bar = ttk.Frame(controls)
        bar.pack(side=tk.TOP, fill=tk.X)
        ttk.Label(bar, text="DT code:").pack(side=tk.LEFT)
        self.dt_var = tk.StringVar(value=dt_string or DEFAULT_DT)
        ttk.Entry(bar, textvariable=self.dt_var, width=58).pack(
            side=tk.LEFT, padx=4, fill=tk.X, expand=True)
        ttk.Button(bar, text="Load / new root", command=self.load).pack(
            side=tk.LEFT, padx=2)
        self.snappy_var = tk.BooleanVar(value=self.use_snappy_global)
        ttk.Checkbutton(bar, text="SnapPy global", variable=self.snappy_var
                        ).pack(side=tk.LEFT, padx=6)
        tk.Button(bar, text="?", width=2, bg="#dbeafe", activebackground="#bfdbfe",
                  relief=tk.RAISED, command=self._show_snappy_help
                  ).pack(side=tk.LEFT, padx=(0, 4))

        # Backtrack-assisted simplification controls (V3.4).
        self.backtrack_var = tk.BooleanVar(value=self.init_backtrack_enabled)
        ttk.Checkbutton(bar, text="Backtrack simplify",
                        variable=self.backtrack_var).pack(side=tk.LEFT, padx=(8, 2))
        ttk.Label(bar, text="rounds").pack(side=tk.LEFT)
        self.backtrack_rounds_var = tk.StringVar(value=str(self.init_backtrack_rounds))
        ttk.Entry(bar, textvariable=self.backtrack_rounds_var, width=5).pack(
            side=tk.LEFT, padx=(1, 4))
        ttk.Label(bar, text="steps").pack(side=tk.LEFT)
        self.backtrack_steps_var = tk.StringVar(value=str(self.init_backtrack_steps))
        ttk.Entry(bar, textvariable=self.backtrack_steps_var, width=5).pack(
            side=tk.LEFT, padx=(1, 4))
        tk.Button(bar, text="?", width=2, bg="#dbeafe", activebackground="#bfdbfe",
                  relief=tk.RAISED, command=self._show_backtrack_help
                  ).pack(side=tk.LEFT, padx=(0, 4))

        order_bar = ttk.Frame(controls)
        order_bar.pack(side=tk.TOP, fill=tk.X, pady=(4, 0))
        ttk.Label(order_bar, text="Crossing labels:").pack(side=tk.LEFT)
        self.crossing_input_var = tk.StringVar(value=self.init_crossing_input)
        ttk.Entry(order_bar, textvariable=self.crossing_input_var, width=66).pack(
            side=tk.LEFT, padx=(4, 8), fill=tk.X, expand=True)
        ttk.Button(order_bar, text="Load drawing session",
                   command=self.load_drawing_session).pack(side=tk.LEFT, padx=2)
        ttk.Button(order_bar, text="Close passage windows",
                   command=self.close_children).pack(side=tk.LEFT, padx=2)

    def load(self):
        self.close_children()
        try:
            crossing_order, crossing_map, _kind = self._current_crossing_options()
            # V4.3: the field may carry a flip vector ('DT: [...], [0,1,...]')
            comps, flips = E.parse_dt_flips(self.dt_var.get())
            note = "Original diagram."
            carry, carry_active = None, False
            if flips is None:
                _refuse_unrealisable(comps)  # round 4: the error window, not a traceback
                # V4.3 (round 3): the realisation a bare root is followed
                # with -- spherogram's reading (dt_converter's mended one for
                # a flagged code it reads illegally); shown only where the
                # composite gate flags a code
                carry, flagged, line = _bare_root_realisation(
                    comps, self.negative_even)
                carry_active = bool(flagged and carry)
                if line is not None:
                    print(line, file=(sys.stderr if line.startswith("[warn]")
                                      else sys.stdout))
                    note = "Original diagram.  %s" % line
            if flips is None:
                start = E.Diagram.from_dt_code(comps,
                                               negative_even=self.negative_even,
                                               carry=carry,
                                               carry_active=carry_active)
            else:
                start = E.Diagram.from_dt_code(comps,
                                               negative_even=self.negative_even,
                                               flips=flips)
            _apply_crossing_display_options(
                start,
                crossing_order,
                crossing_map,
                strict=True,
            )
        except Exception as exc:  # noqa: BLE001
            self._error_window("Could not parse/load DT code:\n\n%s" % exc)
            return
        start.display_source = "original"
        self.step_count = 0
        self.open_step_window(start, note=note, snap=None,
                              passage_log=[], parent=None, is_root=True)

    def load_drawing_session(self):
        try:
            from tkinter import filedialog
            path = filedialog.askopenfilename(
                title="Load drawing session",
                filetypes=[("JSON session", "*.json"), ("All files", "*.*")],
            )
            if not path:
                return
            session = load_drawing_session(path)
            self.drawing_options = normalize_drawing_options(
                session.get("drawing_options"))
            self.drawing_session_path = session.get("path") or path
            if session.get("negative_even"):
                self.negative_even = session["negative_even"]
            if session.get("dt"):
                self.dt_var.set(session["dt"])
            if session.get("crossing_input"):
                self.crossing_input_var.set(session["crossing_input"])
            self.load()
            self._toast(
                self.root,
                "Loaded drawing session:\n%s\n\nLayout: %s"
                % (self.drawing_session_path,
                   self.drawing_options.get("layout", "unknown")))
        except Exception as exc:  # noqa: BLE001
            self._error_window("Could not load drawing session:\n\n%s" % exc)

    def close_children(self):
        for w in list(self.windows):
            try:
                if w["container"] is not self.root:
                    w["container"].destroy()
            except Exception:  # noqa: BLE001
                pass
        self.windows = [w for w in self.windows
                        if w.get("container") is self.root and w["alive"]]
        if hasattr(self, "_root_content") and self._root_content is not None:
            for child in self._root_content.winfo_children():
                child.destroy()
            self.windows = [w for w in self.windows
                            if w["container"] is not self.root]

    def _error_window(self, message):
        tk, ttk = self.tk, self.ttk
        if not hasattr(self, "_root_content") or self._root_content is None:
            self._root_content = ttk.Frame(self.root)
            self._root_content.pack(fill=tk.BOTH, expand=True)
        for child in self._root_content.winfo_children():
            child.destroy()
        ttk.Label(self._root_content, text=message, foreground="#a11",
                  wraplength=700, padding=12).pack(anchor="w")

    # ---- create one step window (root uses the main window, others Toplevel) ---
    def open_step_window(self, diagram, note, snap, passage_log, parent=None,
                         is_root=False):
        tk, ttk = self.tk, self.ttk
        from matplotlib.figure import Figure

        step_index = self.step_count
        self.step_count += 1

        if is_root:
            if not hasattr(self, "_root_content") or self._root_content is None:
                self._root_content = ttk.Frame(self.root)
                self._root_content.pack(fill=tk.BOTH, expand=True)
            for child in self._root_content.winfo_children():
                child.destroy()
            container = self.root
            host = self._root_content
            title = "Step 0  -  original"
        else:
            container = tk.Toplevel(self.root)
            container.title("Step %d  -  %s" % (step_index, note.split(":")[0]))
            container.geometry("1180x760")
            try:
                offx = 60 + 34 * (step_index % 8)
                offy = 60 + 30 * (step_index % 8)
                container.geometry("+%d+%d" % (offx, offy))
            except Exception:  # noqa: BLE001
                pass
            host = ttk.Frame(container)
            host.pack(fill=tk.BOTH, expand=True)
            title = "Step %d" % step_index

        # Optimised layout: a resizable split between the diagram (left) and a
        # scrollable properties panel (right), with a hint bar underneath.
        paned = ttk.PanedWindow(host, orient=tk.HORIZONTAL)
        paned.pack(side=tk.TOP, fill=tk.BOTH, expand=True)

        left = ttk.Frame(paned, padding=4)
        right = ttk.Frame(paned, padding=(6, 4))
        paned.add(left, weight=4)
        paned.add(right, weight=2)

        fig = Figure(figsize=(7.2, 7.2), dpi=100)
        ax = fig.add_subplot(111)

        win: Dict[str, Any] = {"container": container, "alive": True,
                               "click_targets": {}, "diagram": diagram,
                               "passage_log": passage_log,
                               "note": note, "snap": snap,
                               "panel": None, "report": None}

        def nearest_crossing(xdata, ydata):
            targets = win["click_targets"]
            if not targets:
                return None, 1e9
            p = np.array([xdata, ydata])
            cid, best = None, 1e9
            for k, xy in targets.items():
                dloc = float(np.hypot(*(np.asarray(xy) - p)))
                if dloc < best:
                    best, cid = dloc, k
            return cid, best

        def hover_test(xdata, ydata):
            if not win["alive"]:
                return False
            cid, best = nearest_crossing(xdata, ydata)
            return cid is not None and best < win.get("hit_radius",
                                                      SINGLE_CROSSING_RADIUS)

        def click_at_data(xdata, ydata):
            if not win["alive"]:
                return
            cid, best = nearest_crossing(xdata, ydata)
            if cid is None or best >= win.get("hit_radius",
                                              SINGLE_CROSSING_RADIUS):
                return
            bt_rounds, bt_steps = self._current_backtrack()
            try:
                nxt, nnote, nsnap = advance(
                    win["diagram"], cid, negative_even=self.negative_even,
                    use_snappy_global=self.snappy_var.get(),
                    backtrack_rounds=bt_rounds, backtrack_steps=bt_steps)
            except Exception as exc:  # noqa: BLE001
                self._toast(container, "Passage failed: %s" % exc)
                return
            self.open_step_window(nxt, note=nnote, snap=nsnap,
                                  passage_log=win["passage_log"] + ["c%d" % (cid + 1)],
                                  parent=win)

        win["click_at_data"] = click_at_data  # hook for automation/testing
        win["hover_test"] = hover_test
        panel = _make_figure_panel(fig, ax, left, tk, ttk, click_at_data,
                                   hover_test=hover_test,
                                   backend=self.gui_backend)
        win["panel"] = panel

        # properties read-out (with a vertical scrollbar)
        ttk.Label(right, text=title, font=("TkDefaultFont", 10, "bold")
                  ).pack(anchor="w")
        text_wrap = ttk.Frame(right)
        text_wrap.pack(anchor="w", fill=tk.BOTH, expand=True)
        scroll = ttk.Scrollbar(text_wrap, orient=tk.VERTICAL)
        report = tk.Text(text_wrap, width=54, height=32, wrap="word",
                         font=("TkFixedFont", 9), relief="flat",
                         background="#faf9f6", yscrollcommand=scroll.set)
        scroll.config(command=report.yview)
        scroll.pack(side=tk.RIGHT, fill=tk.Y)
        report.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        win["report"] = report

        row = ttk.Frame(right)
        row.pack(anchor="w", pady=4)
        ttk.Button(row, text="Save PNG",
                   command=lambda: self._save_view(fig, step_index, "png")
                   ).grid(row=0, column=0, padx=2)
        ttk.Button(row, text="Save SVG",
                   command=lambda: self._save_view(fig, step_index, "svg")
                   ).grid(row=0, column=1, padx=2)
        ttk.Button(row, text="Simplify",
                   command=lambda: simplify_this()
                   ).grid(row=0, column=2, padx=2)
        if not is_root:
            ttk.Button(row, text="Close window",
                       command=lambda: self._close_window(win)
                       ).grid(row=0, column=3, padx=2)

        def redraw_window():
            # DT labels are shown for any pristine DT drawing (original, direct
            # after-passage, or SnapPy-simplified), unless the drawing session
            # explicitly hides them.
            win["click_targets"] = E.render(
                win["diagram"], ax, show_crossing_ids=True, show_dt_labels=True,
                drawing_options=self._current_drawing_options())
            win["hit_radius"] = _hit_radius(win["click_targets"])
            ax.set_title("click a crossing to pass a strand through it",
                         fontsize=10, color="0.35")
            panel["draw_idle"]()
            report.configure(state="normal")
            report.delete("1.0", tk.END)
            report.insert(
                tk.END,
                "Drawing module: %s\nDisplay: %s\nLayout: %s\nSession: %s\n\n%s"
                % (getattr(D, "__name__", "?"), panel["name"],
                   self.drawing_options.get("layout", "unknown"),
                   self.drawing_session_path or "(V4.0 defaults)",
                   properties_text(
                       win["diagram"], win["note"], win["snap"],
                       win["passage_log"],
                       use_snappy_global=self.snappy_var.get())))
            report.configure(state="disabled")

        def simplify_this():
            if not win["alive"]:
                return
            bt_rounds, bt_steps = self._current_backtrack()
            try:
                simplified, snote, ssnap = simplify_current_diagram(
                    win["diagram"], negative_even=self.negative_even,
                    use_snappy_global=self.snappy_var.get(),
                    backtrack_rounds=bt_rounds, backtrack_steps=bt_steps)
            except Exception as exc:  # noqa: BLE001
                self._toast(container, "Simplify failed: %s" % exc)
                return
            win["diagram"] = simplified
            win["note"] = snote
            win["snap"] = ssnap
            redraw_window()

        # hint bar
        ttk.Label(host, relief="groove", anchor="w", padding=(6, 2),
                  text=("Hover a crossing: the cursor turns into a pointing "
                        "hand where a strand passage can be performed. Click to "
                        "branch a new window.")
                  ).pack(side=tk.BOTTOM, fill=tk.X)

        redraw_window()

        if not is_root:
            container.protocol("WM_DELETE_WINDOW",
                               lambda: self._close_window(win))
        self.windows.append(win)

    def _save_view(self, fig, step_index, ext):
        """Save exactly what is on screen (no tight-bbox crop)."""
        path = "strand_passage_step%d.%s" % (step_index, ext)
        try:
            fig.savefig(path, dpi=160)
        except Exception as exc:  # noqa: BLE001
            self._toast(self.root, "Save failed: %s" % exc)
            return
        self._toast(self.root, "Saved %s" % path)

    def _close_window(self, win):
        win["alive"] = False
        try:
            if win["container"] is not self.root:
                win["container"].destroy()
        except Exception:  # noqa: BLE001
            pass

    def _toast(self, parent, message):
        tk, ttk = self.tk, self.ttk
        top = tk.Toplevel(parent)
        top.title("notice")
        ttk.Label(top, text=message, padding=12, wraplength=420).pack()
        ttk.Button(top, text="OK", command=top.destroy).pack(pady=(0, 8))


def run_gui(dt_string=None, negative_even="over", use_snappy_global=True,
            gui_backend="auto", backtrack_enabled=True,
            backtrack_rounds=DEFAULT_BACKTRACK_ROUNDS,
            backtrack_steps=DEFAULT_BACKTRACK_STEPS,
            crossing_order=None, crossing_map=None, crossing_input=None,
            drawing_options=None, drawing_session_path=None):
    import tkinter as tk
    from tkinter import ttk

    root = tk.Tk()
    D.apply_tk_window_icon(root, tk)
    root.geometry("1240x820")
    StrandPassageApp(root, tk, ttk, dt_string,
                     negative_even=negative_even,
                     use_snappy_global=use_snappy_global,
                     gui_backend=gui_backend,
                     backtrack_enabled=backtrack_enabled,
                     backtrack_rounds=backtrack_rounds,
                     backtrack_steps=backtrack_steps,
                     crossing_order=crossing_order,
                     crossing_map=crossing_map,
                     crossing_input=crossing_input,
                     drawing_options=drawing_options,
                     drawing_session_path=drawing_session_path)
    root.mainloop()


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(
        prog="strand_passage_guiV4_3.py",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        description=(
            "DT Link Toolkit strand passage V%s\n"
            "-----------------------------------\n"
            "Interactive tool (and headless modes) for studying strand passages\n"
            "(crossing changes) on a knot/link given by a signed DT code, with\n"
            "component-colour preservation and an optional SnapPy backend.\n\n"
            "Three modes:\n"
            "  * default            open the interactive GUI (click crossings)\n"
            "  * --nongui           write the two-pass spreadsheet + overview SVG\n"
            "  * --demo IDS         render a headless before/after cascade image\n\n"
            "SnapPy (via 'sage -python') enables simplification, Jones/linking\n"
            "invariants and colour-matching; without Sage those are unavailable.\n\n"
            "A DT code may carry a flip vector naming its realisation:\n"
            "  DT: [(10, 6, 2, 4, 12, 8)], [0,1,0,0,0,1]   (the square knot;\n"
            "  the bare code is the granny).  Bit k belongs to the k-th DT\n"
            "  entry; walking along the visit with the smaller label, 0 = the\n"
            "  other strand crosses from your right to your left (spherogram/\n"
            "  SnapPy convention); complementing every bit is the mirror image."
            % VERSION),
        epilog=(
            "examples:\n"
            "  sage -python strand_passage_guiV4_3.py\n"
            "  sage -python strand_passage_guiV4_3.py --dt \"DT: [(4,6,2)]\"\n"
            "  sage -python strand_passage_guiV4_3.py "
            "--dt \"DT: [(10, 6, 2, 4, 12, 8)], [0,1,0,0,0,1]\"\n"
            "  sage -python strand_passage_guiV4_3.py --drawing-session dt_session.json\n"
            "  sage -python strand_passage_guiV4_3.py --backtrack "
            "--backtrack-rounds 50 --backtrack-steps 25\n"
            "  sage -python strand_passage_guiV4_3.py --nongui \\\n"
            "       --dt \"DT: [(-8,-12,16),(-24,-22,-28,-26),(-10,-14,-2),"
            "(-20,-6,-18,-4)]\" \\\n"
            "       --out results.xlsx --backtrack --backtrack-rounds 50\n"
            "  python3 strand_passage_guiV4_3.py --dt \"DT: [(4,6,2)]\" "
            "--demo 2 1 --out chain.png\n"
            "  python3 strand_passage_guiV4_3.py --gui-backend agg   "
            "# if TkAgg won't load\n"))
    ap.add_argument("--dt", default=None, metavar="STR",
                    help="signed DT code string, e.g. \"DT: [(4,6,2)]\", "
                         "optionally with a flip vector naming the "
                         "realisation, e.g. \"DT: [(10, 6, 2, 4, 12, 8)], "
                         "[0,1,0,0,0,1]\" (default: a built-in 4-component "
                         "example)")
    ap.add_argument("--negative-even", choices=["over", "under"], default=None,
                    help="convention for a negative even DT label: the even "
                         "visit is the 'over' (default) or 'under' strand; "
                         "--nongui rewrites an 'under' code in the 'over' "
                         "convention SnapPy reads (entries negated, the same "
                         "diagram)")
    ap.add_argument("--drawing-session", default=None, metavar="PATH",
                    help="load drawing settings from a JSON session saved by "
                         "draw_dt_original_labelsV5_9.py")
    ap.add_argument("--nongui", action="store_true",
                    help="write the two-pass strand-passage spreadsheet (.xlsx) "
                         "and overview SVG, then exit (needs SnapPy + pandas)")
    ap.add_argument("--demo", nargs="*", type=int, default=None, metavar="ID",
                    help="headless: 0-based crossing ids to click in sequence, "
                         "rendering original + each step to --out")
    ap.add_argument("--out", default=None, metavar="PATH",
                    help="output path (.xlsx for --nongui, image for --demo)")
    ap.add_argument("--no-snappy-global", action="store_true",
                    help="disable the SnapPy simplification/invariant path "
                         "(use only the continuity engine)")
    ap.add_argument("--gui-backend", choices=["auto", "tkagg", "agg"],
                    default="auto",
                    help="Tk drawing backend: auto (default), tkagg, or agg "
                         "(agg = static image fallback if TkAgg won't load)")
    ap.add_argument("--crossing-order", default=None,
                    help="displayed crossing IDs ordered by odd DT labels, "
                         "using the same syntax as draw_dt_original_labels")
    ap.add_argument("--crossing-map", default=None,
                    help="alternative explicit map such as 'c1=1,c7=3'; "
                         "do not combine with --crossing-order")
    ap.add_argument("--crossing-labels", default=None,
                    help="combined crossing-label input: assignment text such "
                         "as 'c1=1,c7=3' is detected as a crossing map; "
                         "otherwise it is detected as a crossing order")
    ap.add_argument("--backtrack", action="store_true",
                    help="(kept for compatibility; backtrack is ON by default "
                         "in V4.0 -- use --no-backtrack to disable)")
    ap.add_argument("--no-backtrack", action="store_true",
                    help="disable backtrack-assisted SnapPy simplification "
                         "(V4.0 enables it by default)")
    ap.add_argument("--backtrack-rounds", type=int, metavar="N",
                    default=DEFAULT_BACKTRACK_ROUNDS,
                    help="backtrack rounds (default %d)"
                         % DEFAULT_BACKTRACK_ROUNDS)
    ap.add_argument("--backtrack-steps", type=int, metavar="K",
                    default=DEFAULT_BACKTRACK_STEPS,
                    help="complication steps per backtrack round (default %d)"
                         % DEFAULT_BACKTRACK_STEPS)
    args = ap.parse_args()

    session = None
    if args.drawing_session:
        try:
            session = load_drawing_session(args.drawing_session)
        except Exception as exc:  # noqa: BLE001
            print("[error] could not load drawing session %s: %s"
                  % (args.drawing_session, exc), file=sys.stderr)
            sys.exit(2)

    drawing_options = normalize_drawing_options(
        session.get("drawing_options") if session else None)
    session_path = session.get("path") if session else None
    session_dt = session.get("dt") if session else None
    session_negative_even = session.get("negative_even") if session else None
    session_crossing_input = session.get("crossing_input") if session else None
    crossing_input = args.crossing_labels
    if crossing_input is None and not (args.crossing_order or args.crossing_map):
        crossing_input = session_crossing_input
    try:
        _split_crossing_display_input(
            crossing_input=crossing_input,
            crossing_order=args.crossing_order,
            crossing_map=args.crossing_map,
        )
    except Exception as exc:  # noqa: BLE001
        print("[error] invalid crossing-label input: %s" % exc, file=sys.stderr)
        sys.exit(2)

    use_snappy_global = not args.no_snappy_global
    negative_even = args.negative_even or session_negative_even or "over"
    backtrack_enabled = not args.no_backtrack           # ON by default (V4.0)
    backtrack_rounds = args.backtrack_rounds if backtrack_enabled else 0
    backtrack_steps = args.backtrack_steps

    # Jones polynomials (and the SnapPy invariant colour-matching that uses them)
    # require Sage; warn once up front when running under plain python.
    warn_if_no_sage()

    if args.nongui or args.demo is not None:
        # a bare code no diagram on the sphere has: --nongui and --demo ended
        # in a traceback ("Link isn't planar", "The crossing graph is not
        # planar"); one line instead
        dt = args.dt or session_dt or (DEFAULT_DT if args.nongui else "DT: [(4,6,2)]")
        why = None if _text_carries_vector(dt) else E.no_diagram_reason(dt)
        if why:
            try:        # round 4b: the code as the dt-flips helper reads it (no '# comment')
                shown = split_dt_flips(dt)[0].strip()
            except ValueError:
                shown = dt
            print("[error] %s: %s, so %s; check the code"
                  % (shown, why, "--nongui has no link to pass strands through"
                     if args.nongui else "--demo has no diagram to draw"),
                  file=sys.stderr)
            sys.exit(2)

    if args.nongui:
        dt = args.dt or session_dt or DEFAULT_DT
        rc = run_nongui(dt, args.out or "strand_passage_results.xlsx",
                        negative_even=negative_even,
                        backtrack_rounds=backtrack_rounds,
                        backtrack_steps=backtrack_steps,
                        crossing_order=args.crossing_order,
                        crossing_map=args.crossing_map,
                        crossing_input=crossing_input,
                        drawing_options=drawing_options,
                        drawing_session_path=session_path)
        sys.exit(rc)

    if args.demo is not None:
        dt = args.dt or session_dt or "DT: [(4,6,2)]"
        out = args.out or "strand_passage_chain_v4_0.png"
        try:
            render_chain(dt, args.demo, out, negative_even=negative_even,
                         use_snappy_global=use_snappy_global,
                         backtrack_rounds=backtrack_rounds,
                         backtrack_steps=backtrack_steps,
                         crossing_order=args.crossing_order,
                         crossing_map=args.crossing_map,
                         crossing_input=crossing_input,
                         drawing_options=drawing_options)
        except ValueError as exc:
            if not _text_carries_vector(dt) and not getattr(exc, "user_facing", False):
                # round 4d (2026-10-20): a bare code the parser cannot read is
                # ONE [error] line, exit 2 (it was a traceback); any other
                # failure of a bare code fails as before
                try:
                    E.parse_dt_flips(dt)
                except ValueError as parse_exc:
                    print("[error] %s" % _unparseable_text(dt, parse_exc), file=sys.stderr)
                    sys.exit(2)
                raise                       # a bare code fails as before
            try:                            # name a parse failure as one
                E.parse_dt_flips(dt)
                what = ""
            except ValueError:
                what = "could not read the DT code: "
            print("[error] %s%s" % (what, exc), file=sys.stderr)
            sys.exit(2)
        print("[ok] wrote", out)
        return

    run_gui(dt_string=args.dt or session_dt, negative_even=negative_even,
            use_snappy_global=use_snappy_global, gui_backend=args.gui_backend,
            backtrack_enabled=backtrack_enabled,
            backtrack_rounds=args.backtrack_rounds,
            backtrack_steps=args.backtrack_steps,
            crossing_order=args.crossing_order,
            crossing_map=args.crossing_map,
            crossing_input=crossing_input,
            drawing_options=drawing_options,
            drawing_session_path=session_path)


if __name__ == "__main__":
    main()
